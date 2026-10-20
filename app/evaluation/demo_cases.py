from __future__ import annotations

import argparse
import json
import os
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook

from app.reports.docx_report import build_docx_report
from app.reports.excel_report import build_excel_report
from app.reports.html_report import build_html_report
from app.reports.package_export import build_review_package_zip
from app.workflows.single_review import run_single_review_workflow


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_MANIFEST_PATH = ROOT / "demo_cases" / "cases_manifest.json"

OFFLINE_ENV_KEYS = {
    "LEGAL_LLM_API_KEY": None,
    "OPENAI_API_KEY": None,
    "RAG_RETRIEVAL_BACKEND": "lexical",
    "RAG_EMBEDDING_PROVIDER": "none",
}


@dataclass
class DemoCaseRun:
    case_id: str
    title: str
    mode: str = ""
    mode_label: str = ""
    issue_count: int = 0
    severity_summary: dict[str, int] = field(default_factory=dict)
    issue_type_summary: dict[str, int] = field(default_factory=dict)
    issue_details: list[dict[str, str]] = field(default_factory=list)
    llm_mode: str | None = None
    fallback_reason: str | None = None
    retrieval_status_summary: dict[str, str] = field(default_factory=dict)
    retrieval_best_docs: dict[str, str | None] = field(default_factory=dict)
    privacy_summary: dict[str, bool] = field(default_factory=dict)
    passed: bool = False
    failures: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "case_id": self.case_id,
            "title": self.title,
            "mode": self.mode,
            "mode_label": self.mode_label,
            "issue_count": self.issue_count,
            "severity_summary": self.severity_summary,
            "issue_type_summary": self.issue_type_summary,
            "issue_details": self.issue_details,
            "llm_mode": self.llm_mode,
            "fallback_reason": self.fallback_reason,
            "retrieval_status_summary": self.retrieval_status_summary,
            "retrieval_best_docs": self.retrieval_best_docs,
            "privacy_summary": self.privacy_summary,
            "passed": self.passed,
            "failures": self.failures,
            "warnings": self.warnings,
            "errors": self.errors,
        }


def load_manifest(path: str | Path = DEFAULT_MANIFEST_PATH) -> dict[str, Any]:
    manifest_path = Path(path)
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict) or not isinstance(payload.get("cases"), list):
        raise ValueError("Demo case manifest must contain a top-level 'cases' list.")
    for case in payload["cases"]:
        _validate_case_shape(case)
    return payload


def list_cases(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    return list(manifest["cases"])


def resolve_case_path(value: str | None) -> Path | None:
    if not value:
        return None
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def run_cases(
    manifest_path: str | Path = DEFAULT_MANIFEST_PATH,
    case_id: str | None = None,
    offline: bool = True,
) -> list[DemoCaseRun]:
    manifest = load_manifest(manifest_path)
    cases = list_cases(manifest)
    ensure_case_files(cases)
    if case_id:
        cases = [case for case in cases if case["case_id"] == case_id]
        if not cases:
            raise ValueError(f"Unknown demo case: {case_id}")
    return [run_case(case, offline=offline) for case in cases]


def run_case(case: dict[str, Any], offline: bool = True) -> DemoCaseRun:
    run = DemoCaseRun(case_id=case["case_id"], title=case.get("title", case["case_id"]))
    env_context = _offline_environment() if offline else _null_context()
    with env_context:
        document_path = resolve_case_path(case["document_path"])
        base_info_path = resolve_case_path(case.get("base_info_path"))
        evidence_index_path = resolve_case_path(case.get("evidence_index_path"))
        if document_path is None:
            run.errors.append("document_path is required")
            run.failures.append("document_path is required")
            return run

        workflow_result = run_single_review_workflow(
            document_path,
            base_info_path=base_info_path,
            evidence_index_path=evidence_index_path,
            document_profile=case.get("document_profile"),
            source_filename=document_path.name,
        )
        run.mode = workflow_result.mode
        run.mode_label = workflow_result.mode_label
        run.warnings = workflow_result.warnings
        run.errors = workflow_result.errors
        if workflow_result.review_result is None:
            run.failures.extend(workflow_result.errors or ["workflow returned no review result"])
            run.passed = False
            return run

        review_result = workflow_result.review_result
        run.issue_count = len(review_result.findings)
        run.severity_summary = dict(Counter(finding.severity for finding in review_result.findings))
        run.issue_type_summary = dict(Counter(str(finding.risk_type) for finding in review_result.findings))
        run.issue_details = [
            {
                "risk_type": str(finding.risk_type),
                "severity": finding.severity,
                "field_name": finding.field_name,
            }
            for finding in review_result.findings[:20]
        ]
        if review_result.structured_extraction:
            run.llm_mode = review_result.structured_extraction.mode
            run.fallback_reason = review_result.structured_extraction.fallback_reason
        run.retrieval_status_summary = {
            status.kb_type: status.status for status in review_result.knowledge_context.retrieval_statuses
        }
        run.retrieval_best_docs = {
            status.kb_type: status.best_doc_id for status in review_result.knowledge_context.retrieval_statuses
        }
        run.privacy_summary = _check_privacy_summary(
            review_result.findings,
            review_result.external_fields,
            workflow_result.mode_label,
            review_result.structured_extraction,
            review_result.knowledge_context,
        )
        run.failures = validate_run(case, run)
        run.passed = not run.failures
        return run


def validate_run(case: dict[str, Any], run: DemoCaseRun) -> list[str]:
    expected = case.get("expected") or {}
    failures: list[str] = []

    if expected.get("mode") and run.mode != expected["mode"]:
        failures.append(f"expected mode {expected['mode']}, got {run.mode}")

    min_issue_count = expected.get("min_issue_count")
    if min_issue_count is not None and run.issue_count < int(min_issue_count):
        failures.append(f"expected at least {min_issue_count} issues, got {run.issue_count}")

    max_issue_count = expected.get("max_issue_count")
    if max_issue_count is not None and run.issue_count > int(max_issue_count):
        failures.append(f"expected at most {max_issue_count} issues, got {run.issue_count}")

    for issue_type in expected.get("expected_issue_types", []):
        if run.issue_type_summary.get(issue_type, 0) <= 0:
            failures.append(f"expected issue type {issue_type}")

    for issue_type in expected.get("expected_absent_issue_types", []):
        if run.issue_type_summary.get(issue_type, 0) > 0:
            failures.append(f"forbidden issue type appeared: {issue_type}")

    if expected.get("max_high_severity") is not None and run.severity_summary.get("high", 0) > int(expected["max_high_severity"]):
        failures.append(f"expected high severity <= {expected['max_high_severity']}, got {run.severity_summary.get('high', 0)}")

    if expected.get("expected_llm_mode") and run.llm_mode != expected["expected_llm_mode"]:
        failures.append(f"expected LLM mode {expected['expected_llm_mode']}, got {run.llm_mode}")

    if expected.get("expected_fallback_reason") and run.fallback_reason != expected["expected_fallback_reason"]:
        failures.append(f"expected fallback_reason {expected['expected_fallback_reason']}, got {run.fallback_reason}")

    for kb_type, status in (expected.get("expected_retrieval_statuses") or {}).items():
        if run.retrieval_status_summary.get(kb_type) != status:
            failures.append(f"expected retrieval status {kb_type}:{status}, got {run.retrieval_status_summary.get(kb_type)}")

    if expected.get("expected_privacy_summary"):
        missing = [name for name, present in run.privacy_summary.items() if not present]
        if missing:
            failures.append(f"privacy summary missing from exports: {', '.join(missing)}")

    return failures


def build_summary(results: Iterable[DemoCaseRun]) -> dict[str, Any]:
    items = list(results)
    passed = sum(1 for item in items if item.passed)
    return {
        "total": len(items),
        "passed": passed,
        "failed": len(items) - passed,
        "cases": [item.to_dict() for item in items],
    }


def format_text_summary(summary: dict[str, Any]) -> str:
    lines = ["Synthetic Demo Case Results", f"Total: {summary['total']} | Passed: {summary['passed']} | Failed: {summary['failed']}"]
    for item in summary["cases"]:
        status = "PASS" if item["passed"] else "FAIL"
        lines.append(
            f"- {item['case_id']} [{status}] mode={item['mode']} issues={item['issue_count']} "
            f"severity={item['severity_summary']} issue_types={item['issue_type_summary']} "
            f"llm={item['llm_mode']} fallback={item['fallback_reason']} retrieval={item['retrieval_status_summary']}"
        )
        for failure in item["failures"]:
            lines.append(f"  * {failure}")
    return "\n".join(lines)


def ensure_case_files(cases: Iterable[dict[str, Any]]) -> None:
    try:
        from scripts.generate_demo_cases import DEMO_CASE_VERSION, VERSION_FILE, generate_demo_case_files
    except Exception as exc:  # pragma: no cover - import should be stable in project checkout
        raise FileNotFoundError(f"Demo case generator could not be loaded: {exc}") from exc
    marker_current = VERSION_FILE.exists() and VERSION_FILE.read_text(encoding="utf-8").strip() == DEMO_CASE_VERSION
    missing = [
        path
        for case in cases
        for path in _case_paths(case)
        if path is not None and not path.exists()
    ]
    if missing or not marker_current:
        generate_demo_case_files()
        missing = [
            path
            for case in cases
            for path in _case_paths(case)
            if path is not None and not path.exists()
        ]
    if not missing:
        return
    missing_text = ", ".join(str(path) for path in missing)
    raise FileNotFoundError(f"Demo case files are missing after generation: {missing_text}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run synthetic Legal Document Review demo/evaluation cases.")
    parser.add_argument("--manifest", default=str(DEFAULT_MANIFEST_PATH), help="Path to cases_manifest.json.")
    parser.add_argument("--case", dest="case_id", help="Run a single case_id.")
    parser.add_argument("--json", action="store_true", help="Emit machine-readable JSON summary.")
    args = parser.parse_args(argv)

    try:
        results = run_cases(args.manifest, case_id=args.case_id)
        summary = build_summary(results)
    except Exception as exc:
        if args.json:
            print(json.dumps({"total": 0, "passed": 0, "failed": 1, "error": str(exc)}, ensure_ascii=False, indent=2))
        else:
            print(f"Demo case runner failed: {exc}")
        return 1

    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    else:
        print(format_text_summary(summary))
    return 0 if summary["failed"] == 0 else 1


def _validate_case_shape(case: Any) -> None:
    if not isinstance(case, dict):
        raise ValueError("Each demo case must be an object.")
    for key in ["case_id", "title", "description", "document_path", "expected"]:
        if key not in case:
            raise ValueError(f"Demo case is missing required key: {key}")
    if not isinstance(case["expected"], dict):
        raise ValueError(f"Demo case {case['case_id']} expected must be an object.")


def _check_privacy_summary(findings, external_fields, mode_label, structured_extraction, knowledge_context) -> dict[str, bool]:
    html = build_html_report(
        findings,
        external_fields,
        [],
        mode_label=mode_label,
        structured_extraction=structured_extraction,
        knowledge_context=knowledge_context,
    )
    excel = build_excel_report(
        findings,
        external_fields,
        [],
        structured_extraction=structured_extraction,
        knowledge_context=knowledge_context,
    )
    docx = build_docx_report(
        findings,
        external_fields,
        [],
        mode_label=mode_label,
        structured_extraction=structured_extraction,
        knowledge_context=knowledge_context,
    )
    package = build_review_package_zip(
        findings,
        external_fields,
        [],
        mode_label=mode_label,
        structured_extraction=structured_extraction,
        knowledge_context=knowledge_context,
    )

    excel_workbook = load_workbook(BytesIO(excel), read_only=True, data_only=True)
    try:
        excel_has_summary = "Privacy Summary" in excel_workbook.sheetnames
    finally:
        excel_workbook.close()

    with ZipFile(BytesIO(package)) as archive:
        names = set(archive.namelist())
        zip_has_summary = "privacy_summary.json" in names and "This is a synthetic local demo." in archive.read("privacy_summary.json").decode("utf-8")

    return {
        "html": "This is a synthetic local demo." in html,
        "excel": excel_has_summary,
        "docx": _docx_contains_privacy_summary(docx),
        "zip": zip_has_summary,
    }


def _case_paths(case: dict[str, Any]) -> list[Path | None]:
    return [
        resolve_case_path(case.get("document_path")),
        resolve_case_path(case.get("base_info_path")),
        resolve_case_path(case.get("evidence_index_path")),
    ]


def _docx_contains_privacy_summary(data: bytes) -> bool:
    document = Document(BytesIO(data))
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    return "Privacy Summary" in text and "This is a synthetic local demo." in text


@contextmanager
def _offline_environment():
    previous = {key: os.environ.get(key) for key in OFFLINE_ENV_KEYS}
    try:
        for key, value in OFFLINE_ENV_KEYS.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        yield
    finally:
        for key, value in previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


@contextmanager
def _null_context():
    yield

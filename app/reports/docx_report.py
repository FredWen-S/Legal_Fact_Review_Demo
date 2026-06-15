from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.models.schemas import ExtractedReviewField, ManualReviewRecord, ReviewFinding
from app.reports.excel_report import (
    external_fields_to_dataframe,
    findings_to_dataframe,
    manual_records_to_dataframe,
)
from app.utils import mask_sensitive


DISCLAIMER = "本报告仅基于上传文书及用户提供资料生成，用于辅助人工复核，不构成法律意见；外部来源仍需人工核对，最终判断应由律师或相关人员完成。"


def build_docx_report(
    findings: list[ReviewFinding],
    external_fields: list[ExtractedReviewField] | None = None,
    manual_records: list[ManualReviewRecord] | None = None,
    review_id: str | None = None,
    mode_label: str = "完整复核模式",
    generated_at: datetime | None = None,
    config_version: str | None = None,
    document_profile: str | None = None,
    config_updated_at: str | None = None,
) -> bytes:
    generated_at = generated_at or datetime.now()
    external_fields = external_fields or []
    manual_records = manual_records or []

    document = Document()
    _configure_styles(document)
    title = document.add_heading("文书复核摘要", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    high = sum(1 for finding in findings if finding.severity == "high")
    medium = sum(1 for finding in findings if finding.severity == "medium")
    low = sum(1 for finding in findings if finding.severity == "low")

    document.add_heading("基本信息", level=1)
    _add_key_value_table(
        document,
        [
            ("review_id", review_id or "未记录"),
            ("生成时间", generated_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("当前模式", mode_label),
            ("文书类型模板", document_profile or "未记录"),
            ("规则配置版本", config_version or "未记录"),
            ("规则配置更新时间", config_updated_at or "未记录"),
            ("findings 总数", str(len(findings))),
            ("高 / 中 / 低风险数量", f"{high} / {medium} / {low}"),
            ("待外部复核字段数量", str(len(external_fields))),
            ("人工复核记录数量", str(len(manual_records))),
        ],
    )

    document.add_heading("免责声明", level=1)
    document.add_paragraph(DISCLAIMER)

    document.add_heading("复核发现汇总", level=1)
    _add_dataframe_table(document, findings_to_dataframe(findings), empty_text="暂无复核发现")

    document.add_heading("待外部复核字段清单", level=1)
    _add_dataframe_table(document, external_fields_to_dataframe(external_fields), empty_text="暂无待外部复核字段")

    document.add_heading("人工复核记录", level=1)
    _add_dataframe_table(document, manual_records_to_dataframe(manual_records), empty_text="暂无人工复核记录")

    document.add_heading("高风险优先处理清单", level=1)
    high_df = findings_to_dataframe([finding for finding in findings if finding.severity == "high"])
    _add_dataframe_table(document, high_df, empty_text="暂无高风险复核发现")

    document.add_heading("需修改字段清单", level=1)
    modify_records = [record for record in manual_records if record.review_status == "需修改"]
    modify_df = pd.DataFrame(
        [
            {
                "字段类型": record.field_type,
                "原提取值": mask_sensitive(record.extracted_value, record.field_type),
                "人工确认值": mask_sensitive(record.confirmed_value, record.field_type),
                "出现位置": record.source_location,
                "复核备注": record.reviewer_note,
            }
            for record in modify_records
        ],
        columns=["字段类型", "原提取值", "人工确认值", "出现位置", "复核备注"],
    )
    _add_dataframe_table(document, modify_df, empty_text="暂无需修改字段")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)


def _add_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "项目"
    header[1].text = "内容"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value


def _add_dataframe_table(document: Document, df: pd.DataFrame, empty_text: str) -> None:
    if df.empty:
        document.add_paragraph(empty_text)
        return
    table = document.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, column in enumerate(df.columns):
        header_cells[index].text = str(column)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(df.columns):
            value = row[column]
            cells[index].text = "" if pd.isna(value) else str(value)

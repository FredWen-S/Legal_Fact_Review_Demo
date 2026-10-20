from __future__ import annotations

from pathlib import Path

from docx import Document

from app.errors import EmptyFileError, ParseError
from app.models.schemas import TextChunk


def parse_docx(path: str | Path) -> list[TextChunk]:
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ParseError("DOCX 解析失败。请确认上传的是未损坏的 .docx 文件，且不是加密文档。") from exc
    chunks: list[TextChunk] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            chunks.append(TextChunk(text=text, location=f"DOCX 段落 {index}"))

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(
                    TextChunk(
                        text=" | ".join(cells),
                        location=f"DOCX 表格 {table_index} 行 {row_index}",
                    )
                )
    if not chunks:
        raise EmptyFileError("DOCX 未解析到正文内容。请确认文书不是空文件，或内容不全在图片中。")
    return chunks

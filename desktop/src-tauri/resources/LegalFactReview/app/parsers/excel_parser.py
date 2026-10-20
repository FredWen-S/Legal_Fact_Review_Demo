from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.errors import EmptyFileError, ExcelSchemaError, ParseError
from app.models.schemas import BaseInfoItem, EvidenceItem


def _header_map(row: tuple[object, ...]) -> dict[str, int]:
    return {str(cell).strip(): index for index, cell in enumerate(row) if cell is not None}


def _load_rows(path: str | Path, label: str) -> list[tuple[object, ...]]:
    try:
        workbook = load_workbook(path, data_only=True)
    except Exception as exc:
        raise ParseError(f"{label}解析失败。请确认上传的是未损坏的 .xlsx 文件。") from exc
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise EmptyFileError(f"{label}为空。请上传包含表头和数据行的 Excel 文件。")
    return rows


def _require_headers(headers: dict[str, int], required: set[str], label: str) -> None:
    missing = sorted(required - set(headers))
    if missing:
        raise ExcelSchemaError(f"{label}缺少必要列：{', '.join(missing)}。请按模板补齐表头后重新上传。")


def parse_base_info(path: str | Path) -> list[BaseInfoItem]:
    rows = _load_rows(path, "基础信息表")
    headers = _header_map(rows[0])
    _require_headers(headers, {"字段名", "标准值", "更新时间", "备注"}, "基础信息表")
    result: list[BaseInfoItem] = []
    for row in rows[1:]:
        field = row[headers.get("字段名", -1)] if "字段名" in headers else None
        value = row[headers.get("标准值", -1)] if "标准值" in headers else None
        if field is None or value is None:
            continue
        updated_at = row[headers.get("更新时间", -1)] if "更新时间" in headers else None
        remark = row[headers.get("备注", -1)] if "备注" in headers else None
        result.append(
            BaseInfoItem(
                field_name=str(field).strip(),
                standard_value=str(value).strip(),
                updated_at=str(updated_at).strip() if updated_at else None,
                remark=str(remark).strip() if remark else None,
            )
        )
    if not result:
        raise EmptyFileError("基础信息表没有有效数据行。请至少填写“字段名”和“标准值”。")
    return result


def parse_evidence_index(path: str | Path) -> list[EvidenceItem]:
    rows = _load_rows(path, "证据目录")
    headers = _header_map(rows[0])
    _require_headers(headers, {"证据编号", "证据名称", "页码/附件位置", "备注"}, "证据目录")
    result: list[EvidenceItem] = []
    for row in rows[1:]:
        no = row[headers.get("证据编号", -1)] if "证据编号" in headers else None
        name = row[headers.get("证据名称", -1)] if "证据名称" in headers else None
        if no is None:
            continue
        location = row[headers.get("页码/附件位置", -1)] if "页码/附件位置" in headers else None
        remark = row[headers.get("备注", -1)] if "备注" in headers else None
        result.append(
            EvidenceItem(
                evidence_no=str(no).strip(),
                evidence_name=str(name).strip() if name else "",
                location=str(location).strip() if location else None,
                remark=str(remark).strip() if remark else None,
            )
        )
    if not result:
        raise EmptyFileError("证据目录没有有效数据行。请至少填写“证据编号”。")
    return result

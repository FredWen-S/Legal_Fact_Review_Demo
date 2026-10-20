from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from app.models.schemas import ExtractedReviewField, ManualReviewRecord, ReviewFinding
from app.utils import mask_sensitive


ANNOTATION_DISCLAIMER = "本复核提示版文书仅用于辅助人工复核，不构成法律意见，不代表系统已确认文书存在错误，最终判断应由律师或相关人员完成。"


@dataclass(frozen=True)
class _Prompt:
    key: str
    source_location: str
    document_value: str | None
    text: str
    severity: str = ""


def annotated_docx_filename(original_filename: str) -> str:
    path = Path(original_filename)
    stem = path.stem or "文书"
    return f"{stem}_复核提示版.docx"


def build_annotated_docx(
    source_docx: str | Path | bytes,
    findings: list[ReviewFinding],
    external_fields: list[ExtractedReviewField] | None = None,
    manual_records: list[ManualReviewRecord] | None = None,
) -> bytes:
    document = _open_document(source_docx)
    external_fields = external_fields or []
    manual_records = manual_records or []
    prompts = _build_prompts(findings, external_fields, manual_records)

    unlocated: list[_Prompt] = []
    for prompt in prompts:
        target = _find_target_paragraph(document, prompt)
        if target is None:
            unlocated.append(prompt)
            continue
        _insert_prompt_after(target, prompt.text)

    _append_summary(document, findings, external_fields, manual_records, unlocated)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _open_document(source_docx: str | Path | bytes) -> DocxDocument:
    if isinstance(source_docx, bytes):
        return Document(BytesIO(source_docx))
    return Document(str(source_docx))


def _build_prompts(
    findings: list[ReviewFinding],
    external_fields: list[ExtractedReviewField],
    manual_records: list[ManualReviewRecord],
) -> list[_Prompt]:
    manual_by_field = {record.field_id: record for record in manual_records}
    prompts: list[_Prompt] = []
    for finding in findings:
        prompts.append(
            _Prompt(
                key=f"finding:{finding.risk_id}",
                source_location=finding.source_location,
                document_value=finding.document_value,
                severity=finding.severity,
                text=_finding_prompt_text(finding, _match_manual_for_finding(finding, manual_records)),
            )
        )
    for field in external_fields:
        manual = manual_by_field.get(field.field_id)
        prompts.append(
            _Prompt(
                key=f"external:{field.field_id}",
                source_location=field.source_location,
                document_value=field.value,
                text=_external_field_prompt_text(field, manual),
            )
        )
    return prompts


def _finding_prompt_text(finding: ReviewFinding, manual: ManualReviewRecord | None) -> str:
    lines = [
        "【复核提示】",
        f"风险等级：{finding.severity}",
        f"风险类型：{finding.risk_type}",
        f"字段：{finding.field_name}",
        f"文书值：{mask_sensitive(finding.document_value, finding.field_name) or ''}",
    ]
    if finding.reference_value:
        lines.append(f"参考值：{mask_sensitive(finding.reference_value, finding.field_name) or ''}")
    lines.extend(
        [
            f"原因：{finding.reason}",
            f"建议动作：{finding.suggested_action}",
        ]
    )
    if manual:
        lines.append(f"复核状态：{manual.review_status}")
        if manual.confirmed_value:
            lines.append(f"人工确认值：{mask_sensitive(manual.confirmed_value, manual.field_type) or ''}")
    return "\n".join(lines)


def _external_field_prompt_text(field: ExtractedReviewField, manual: ManualReviewRecord | None) -> str:
    lines = [
        "【复核提示】",
        "风险等级：待人工复核",
        "风险类型：EXTERNAL_REVIEW_FIELD",
        f"字段：{field.field_type}",
        f"文书值：{mask_sensitive(field.value, field.field_type) or ''}",
        f"原因：该字段已被抽取为待外部复核字段。",
        f"建议动作：{field.suggested_action}",
    ]
    if manual:
        lines.append(f"复核状态：{manual.review_status}")
        if manual.confirmed_value:
            lines.append(f"人工确认值：{mask_sensitive(manual.confirmed_value, manual.field_type) or ''}")
    return "\n".join(lines)


def _match_manual_for_finding(finding: ReviewFinding, records: list[ManualReviewRecord]) -> ManualReviewRecord | None:
    for record in records:
        if record.field_type == finding.field_name and record.source_location == finding.source_location:
            return record
    return None


def _find_target_paragraph(document: DocxDocument, prompt: _Prompt) -> Paragraph | None:
    value = (prompt.document_value or "").strip()
    if value and "*" not in value:
        for paragraph in _iter_all_paragraphs(document):
            if value in paragraph.text:
                return paragraph

    paragraph_index = _paragraph_index_from_location(prompt.source_location)
    if paragraph_index is not None:
        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        if 1 <= paragraph_index <= len(paragraphs):
            return paragraphs[paragraph_index - 1]

    return None


def _iter_all_paragraphs(document: DocxDocument):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _paragraph_index_from_location(location: str) -> int | None:
    match = re.search(r"(?:DOCX\s*)?段落\s*(\d+)", location or "")
    if match:
        return int(match.group(1))
    return None


def _insert_prompt_after(paragraph: Paragraph, text: str) -> None:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    _fill_prompt_paragraph(new_paragraph, text)


def _fill_prompt_paragraph(paragraph: Paragraph, text: str) -> None:
    lines = text.splitlines()
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        if index == 0:
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _append_summary(
    document: DocxDocument,
    findings: list[ReviewFinding],
    external_fields: list[ExtractedReviewField],
    manual_records: list[ManualReviewRecord],
    unlocated: list[_Prompt],
) -> None:
    document.add_page_break()
    document.add_heading("复核提示汇总", level=1)

    document.add_heading("高风险优先处理清单", level=2)
    high_findings = [finding for finding in findings if finding.severity == "high"]
    if high_findings:
        for finding in high_findings:
            document.add_paragraph(_finding_prompt_text(finding, _match_manual_for_finding(finding, manual_records)))
    else:
        document.add_paragraph("暂无高风险复核提示。")

    document.add_heading("未定位复核提示", level=2)
    if unlocated:
        for prompt in unlocated:
            document.add_paragraph(prompt.text)
    else:
        document.add_paragraph("暂无未定位复核提示。")

    document.add_heading("待外部复核字段", level=2)
    if external_fields:
        for field in external_fields:
            document.add_paragraph(_external_field_prompt_text(field, _manual_for_field(field, manual_records)))
    else:
        document.add_paragraph("暂无待外部复核字段。")

    document.add_heading("人工复核记录", level=2)
    if manual_records:
        for record in manual_records:
            text = (
                f"字段：{record.field_type}\n"
                f"提取值：{mask_sensitive(record.extracted_value, record.field_type) or ''}\n"
                f"复核状态：{record.review_status}\n"
                f"人工确认值：{mask_sensitive(record.confirmed_value, record.field_type) or ''}\n"
                f"复核备注：{record.reviewer_note or ''}"
            )
            document.add_paragraph(text)
    else:
        document.add_paragraph("暂无人工复核记录。")

    document.add_heading("免责声明", level=2)
    document.add_paragraph(ANNOTATION_DISCLAIMER)


def _manual_for_field(field: ExtractedReviewField, records: list[ManualReviewRecord]) -> ManualReviewRecord | None:
    for record in records:
        if record.field_id == field.field_id:
            return record
    return None

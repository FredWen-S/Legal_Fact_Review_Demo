from __future__ import annotations

from pathlib import Path
import sys
import json
import sqlite3
from datetime import datetime
from tempfile import TemporaryDirectory
from io import BytesIO
from zipfile import ZipFile
from docx import Document
import fitz

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.parsers.docx_parser import parse_docx
from app.parsers.excel_parser import parse_base_info, parse_evidence_index
from app.extractors.external_review_fields import build_external_review_fields
from app.reports.excel_report import build_excel_report, external_fields_to_dataframe, findings_to_dataframe
from app.reports.html_report import build_html_report
from app.reports.docx_report import build_docx_report
from app.reports.package_export import build_review_package_zip
from app.manual_review import default_manual_records
from app.storage.sqlite_store import fetch_manual_review_records, save_manual_review_records
from app.review_engine import review_document, review_single_document
from app.errors import OcrDependencyError
from app.parsers.pdf_parser import is_scanned_pdf
from app.parsers.ocr_pdf_parser import parse_pdf_with_ocr
from scripts.build_portable import build_portable
from scripts.build_release_zip import build_release_zip
from scripts.generate_test_files import TEST_CASES_DIR, generate_test_files
from scripts.windows.preflight_check import run_preflight
from app.reports.case_package import build_case_review_package_zip
from app.reports.docx_annotation import ANNOTATION_DISCLAIMER, build_annotated_docx
from app.config.review_config import default_config, export_config, import_config_file, load_review_config, load_review_config_with_status, restore_default_config
from app.config.rule_config_service import import_rules_with_history
from app.models.schemas import TextChunk
from app.storage.sqlite_store import create_case_workspace, init_db, list_review_runs
from app.workflows.batch_review import BatchInputFile, run_batch_review
from app.workflows.case_differences import build_case_field_differences
from app.extractors.field_extractor import extract_fields


SAMPLE_DIR = ROOT / "app" / "sample_data"


def main() -> None:
    document = SAMPLE_DIR / "sample_document.docx"
    base_info = SAMPLE_DIR / "sample_base_info.xlsx"
    evidence = SAMPLE_DIR / "sample_evidence_index.xlsx"

    missing = [path for path in [document, base_info, evidence] if not path.exists()]
    if missing:
        raise SystemExit(f"示例数据缺失：{', '.join(str(path) for path in missing)}")

    chunks = parse_docx(document)
    external_fields = build_external_review_fields(chunks)
    findings = review_document(chunks, parse_base_info(base_info), parse_evidence_index(evidence))
    risk_types = {finding.risk_type for finding in findings}

    required_risks = {
        "PLACEHOLDER_REMAINING",
        "INTERNAL_CONFLICT",
        "SOURCE_MISMATCH",
        "EVIDENCE_REFERENCE_MISSING",
        "AMOUNT_INCONSISTENCY",
        "DATE_LOGIC_WARNING",
        "FORMAT_ERROR",
    }
    missing_risks = sorted(required_risks - risk_types)
    if missing_risks:
        raise SystemExit(f"示例复核未覆盖预期风险类型：{', '.join(missing_risks)}")
    if len(findings) < 5:
        raise SystemExit(f"示例复核发现数量过少：{len(findings)}")
    if any(not finding.source_location.strip() for finding in findings):
        raise SystemExit("存在缺少 source_location 的 finding")
    if any("建议人工复核" not in finding.suggested_action for finding in findings):
        raise SystemExit("存在未按人工复核表述输出的 suggested_action")

    manual_records = default_manual_records("acceptance-full", external_fields)
    manual_records[0] = manual_records[0].model_copy(
        update={
            "review_status": "已确认",
            "review_source": "内部客户库",
            "reviewer_note": "验收脚本生成的人工复核记录",
            "reviewed_at": datetime.now(),
        }
    )
    excel_bytes = build_excel_report(findings, external_fields, manual_records)
    html = build_html_report(findings, external_fields, manual_records, mode_label="完整复核模式")
    if excel_bytes[:2] != b"PK":
        raise SystemExit("Excel 报告生成异常")
    if "不构成法律意见" not in html:
        raise SystemExit("HTML 报告缺少免责声明")
    if "人工复核记录" not in html:
        raise SystemExit("HTML 报告缺少人工复核记录")

    df = findings_to_dataframe(findings)
    if "91310000MA1K000001" in df.to_string():
        raise SystemExit("复核表格中存在未脱敏统一社会信用代码")

    single_findings = review_single_document(chunks)
    single_risk_types = {finding.risk_type for finding in single_findings}
    if "SOURCE_MISMATCH" in single_risk_types:
        raise SystemExit("单文书模式不应生成 SOURCE_MISMATCH")
    if "EVIDENCE_REFERENCE_MISSING" in single_risk_types:
        raise SystemExit("单文书模式不应因缺少证据目录生成 EVIDENCE_REFERENCE_MISSING")
    if len(external_fields) < 5:
        raise SystemExit(f"单文书模式待外部复核字段过少：{len(external_fields)}")
    if any(not field.source_location.strip() for field in external_fields):
        raise SystemExit("存在缺少 source_location 的待外部复核字段")
    if any("建议人工复核" not in field.suggested_action for field in external_fields):
        raise SystemExit("存在未按人工复核表述输出的待外部复核字段")

    with TemporaryDirectory() as tmpdir:
        db_path = Path(tmpdir) / "acceptance.sqlite"
        save_manual_review_records(manual_records, db_path)
        loaded_records = fetch_manual_review_records("acceptance-full", db_path)
        if not loaded_records:
            raise SystemExit("人工复核记录未能保存到 SQLite")
        if not any(record.review_status == "已确认" for record in loaded_records):
            raise SystemExit("人工复核记录读取结果异常")

    single_manual_records = default_manual_records("acceptance-single", external_fields)
    single_manual_records[0] = single_manual_records[0].model_copy(
        update={
            "review_status": "需修改",
            "confirmed_value": "人工确认值示例",
            "review_source": "其他",
            "reviewer_note": "验收脚本生成的单文书人工复核记录",
            "reviewed_at": datetime.now(),
        }
    )
    single_excel = build_excel_report(single_findings, external_fields, single_manual_records)
    single_html = build_html_report(single_findings, external_fields, single_manual_records, mode_label="单文书字段抽取模式")
    if single_excel[:2] != b"PK":
        raise SystemExit("单文书模式 Excel 报告生成异常")
    if "待外部复核字段清单" not in single_html:
        raise SystemExit("单文书模式 HTML 报告缺少字段清单")
    if "人工复核记录" not in single_html:
        raise SystemExit("单文书模式 HTML 报告缺少人工复核记录")
    external_df = external_fields_to_dataframe(external_fields)
    if "91310000MA1K000001" in external_df.to_string():
        raise SystemExit("待外部复核字段中存在未脱敏统一社会信用代码")

    docx_bytes = build_docx_report(findings, external_fields, manual_records, review_id="acceptance-full", mode_label="完整复核模式")
    docx_text = _extract_docx_text(docx_bytes)
    if "本报告仅基于上传文书及用户提供资料生成" not in docx_text:
        raise SystemExit("Word 复核摘要缺少免责声明")
    if "复核发现汇总" not in docx_text or "人工复核记录" not in docx_text:
        raise SystemExit("Word 复核摘要缺少关键章节")
    if "110101199001011234" in docx_text or "13800138000" in docx_text:
        raise SystemExit("Word 复核摘要存在未脱敏身份证号或手机号")

    zip_bytes = build_review_package_zip(findings, external_fields, manual_records, review_id="acceptance-full", mode_label="完整复核模式")
    with ZipFile(BytesIO(zip_bytes)) as package:
        names = set(package.namelist())
    expected = {"复核工作底稿.xlsx", "复核报告.html", "文书复核摘要.docx"}
    if not expected.issubset(names):
        raise SystemExit(f"复核包缺少文件：{', '.join(sorted(expected - names))}")

    original_docx_bytes = document.read_bytes()
    unlocated_finding = findings[0].model_copy(
        update={
            "risk_id": "acceptance-unlocated",
            "document_value": "不存在的定位值",
            "source_location": "文书正文",
        }
    )
    annotated_bytes = build_annotated_docx(document, findings + [unlocated_finding], external_fields, manual_records)
    annotated_text = _extract_docx_text(annotated_bytes)
    if document.read_bytes() != original_docx_bytes:
        raise SystemExit("Word 复核提示版生成覆盖了原始 DOCX")
    if "【复核提示】" not in annotated_text or "建议人工复核" not in annotated_text:
        raise SystemExit("Word 复核提示版缺少复核提示内容")
    if "未定位复核提示" not in annotated_text or "不存在的定位值" not in annotated_text:
        raise SystemExit("Word 复核提示版未将定位失败项写入文末汇总")
    if ANNOTATION_DISCLAIMER not in annotated_text:
        raise SystemExit("Word 复核提示版缺少免责声明")

    generated_test_files = generate_test_files()
    if len(generated_test_files) < 20:
        raise SystemExit(f"测试文件生成数量不足：{len(generated_test_files)}")
    case_dirs = sorted(path for path in TEST_CASES_DIR.iterdir() if path.is_dir())
    if len(case_dirs) < 6:
        raise SystemExit(f"测试 case 目录数量不足：{len(case_dirs)}")
    if any(not (case_dir / "README.md").exists() for case_dir in case_dirs):
        raise SystemExit("存在缺少 README.md 的测试 case")

    case01 = TEST_CASES_DIR / "case_01_full_review_basic"
    case01_findings = review_document(
        parse_docx(case01 / "sample_complaint.docx"),
        parse_base_info(case01 / "base_info.xlsx"),
        parse_evidence_index(case01 / "evidence_index.xlsx"),
    )
    case01_risks = {finding.risk_type for finding in case01_findings}
    expected_case01_risks = {
        "SOURCE_MISMATCH",
        "AMOUNT_INCONSISTENCY",
        "EVIDENCE_REFERENCE_MISSING",
        "PLACEHOLDER_REMAINING",
        "FORMAT_ERROR",
        "DATE_LOGIC_WARNING",
    }
    missing_case01 = expected_case01_risks - case01_risks
    if missing_case01:
        raise SystemExit(f"Case 01 未覆盖预期风险类型：{', '.join(sorted(missing_case01))}")

    case02 = TEST_CASES_DIR / "case_02_single_document_only"
    case02_chunks = parse_docx(case02 / "single_document.docx")
    case02_findings = review_single_document(case02_chunks)
    case02_risks = {finding.risk_type for finding in case02_findings}
    if "SOURCE_MISMATCH" in case02_risks or "EVIDENCE_REFERENCE_MISSING" in case02_risks:
        raise SystemExit("Case 02 单文书模式出现外部比对或证据目录误报")
    if not build_external_review_fields(case02_chunks):
        raise SystemExit("Case 02 未生成待外部复核字段")

    case06 = TEST_CASES_DIR / "case_06_clean_document"
    case06_findings = review_document(
        parse_docx(case06 / "clean_document.docx"),
        parse_base_info(case06 / "base_info.xlsx"),
        parse_evidence_index(case06 / "evidence_index.xlsx"),
    )
    case06_risks = {finding.risk_type for finding in case06_findings}
    if len(case06_findings) > 1 or {"SOURCE_MISMATCH", "EVIDENCE_REFERENCE_MISSING", "PLACEHOLDER_REMAINING"} & case06_risks:
        raise SystemExit(f"Case 06 干净场景风险过多或存在误报：{sorted(case06_risks)}")

    with TemporaryDirectory() as tmpdir:
        config_root = Path(tmpdir)
        default_config_path = config_root / "review_rules.yml"
        default_config_path.write_text(export_config(default_config()), encoding="utf-8")
        loaded_default = load_review_config(default_config_path)
        if "civil_complaint" not in loaded_default.document_profiles:
            raise SystemExit("默认配置未加载文书类型模板")
        exported = export_config(loaded_default)
        if "document_profiles" not in exported or "field_dictionary" not in exported:
            raise SystemExit("配置导出缺少文书类型模板或字段词典")
        broken_path = config_root / "broken.yml"
        broken_path.write_text("enabled_risk_types: [NOT_A_RISK]\n", encoding="utf-8")
        if not load_review_config_with_status(broken_path).fallback_used:
            raise SystemExit("配置损坏时未 fallback 默认配置")
        case09 = TEST_CASES_DIR / "case_09_custom_rules_profile"
        custom_config = import_rules_with_history(
            case09 / "custom_review_rules.yml",
            target_path=default_config_path,
            db_path=config_root / "config_history.sqlite",
            change_note="acceptance case09",
        )
        if custom_config.config_version != "case09-custom-v1":
            raise SystemExit("Case 09 自定义配置导入失败")
        case09_chunks = parse_docx(case09 / "lawyer_letter.docx")
        case09_findings = review_single_document(case09_chunks, config=custom_config)
        if "PLACEHOLDER_REMAINING" in {finding.risk_type for finding in case09_findings}:
            raise SystemExit("Case 09 禁用 PLACEHOLDER_REMAINING 后仍输出占位符风险")
        json_data = default_config().model_dump(mode="json")
        json_data["config_version"] = "acceptance-json-v1"
        json_path = config_root / "review_rules.json"
        json_path.write_text(json.dumps(json_data, ensure_ascii=False), encoding="utf-8")
        imported_json, _ = import_config_file(json_path, target_path=default_config_path, backup_dir=config_root / "backups")
        if imported_json.config_version != "acceptance-json-v1":
            raise SystemExit("JSON 规则配置导入失败")
        before_bad_import = default_config_path.read_text(encoding="utf-8")
        bad_path = config_root / "bad_rules.yml"
        bad_path.write_text("enabled_risk_types: [BAD_RISK]\n", encoding="utf-8")
        try:
            import_config_file(bad_path, target_path=default_config_path, backup_dir=config_root / "backups")
        except Exception:
            pass
        else:
            raise SystemExit("非法规则配置导入未失败")
        if default_config_path.read_text(encoding="utf-8") != before_bad_import:
            raise SystemExit("非法规则配置导入覆盖了旧配置")
        alias_config = default_config()
        updated_legal_rep = alias_config.field_dictionary["legal_representative"].model_copy(update={"aliases": ["负责人"]})
        alias_config = alias_config.model_copy(update={"field_dictionary": {**alias_config.field_dictionary, "legal_representative": updated_legal_rep}})
        alias_items = extract_fields([TextChunk(text="负责人：张明远。", location="验收段落")], config=alias_config)
        if ("法定代表人", "张明远") not in {(item.field_name, item.value) for item in alias_items}:
            raise SystemExit("字段 alias 未影响字段抽取")
        restore_default_config(target_path=default_config_path, backup_dir=config_root / "backups")
        restored = load_review_config(default_config_path)
        if "PLACEHOLDER_REMAINING" not in restored.enabled_risk_types:
            raise SystemExit("恢复默认配置失败")

    with TemporaryDirectory() as tmpdir:
        old_db = Path(tmpdir) / "old.sqlite"
        conn = sqlite3.connect(old_db)
        try:
            conn.execute(
                """
                CREATE TABLE review_runs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at TEXT NOT NULL,
                    finding_count INTEGER NOT NULL,
                    findings_json TEXT NOT NULL
                )
                """
            )
            conn.execute("INSERT INTO review_runs (created_at, finding_count, findings_json) VALUES ('2026-01-01T00:00:00', 0, '[]')")
            conn.execute(
                """
                CREATE TABLE manual_review_records (
                    record_id TEXT PRIMARY KEY,
                    review_id TEXT NOT NULL,
                    field_id TEXT NOT NULL,
                    field_type TEXT NOT NULL,
                    extracted_value TEXT NOT NULL,
                    confirmed_value TEXT,
                    review_status TEXT NOT NULL,
                    review_source TEXT,
                    reviewer_note TEXT,
                    source_location TEXT NOT NULL,
                    reviewed_at TEXT
                )
                """
            )
            conn.commit()
        finally:
            conn.close()
        init_db(old_db).close()
        init_db(old_db).close()
        conn = sqlite3.connect(old_db)
        try:
            columns = {row[1] for row in conn.execute("PRAGMA table_info(review_runs)").fetchall()}
            tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
            row_count = conn.execute("SELECT COUNT(*) FROM review_runs").fetchone()[0]
        finally:
            conn.close()
        if not {"config_version", "document_profile"}.issubset(columns) or "rule_config_history" not in tables or row_count != 1:
            raise SystemExit("旧 SQLite migration 未正确补列、建表或保留旧记录")

    case08 = TEST_CASES_DIR / "case_08_case_workspace_batch"
    if not case08.exists():
        raise SystemExit("Case 08 测试目录未生成")
    with TemporaryDirectory() as tmpdir:
        workspace_db = Path(tmpdir) / "case_workspace.sqlite"
        case = create_case_workspace(
            "Case 08 批量复核验收案件",
            internal_case_no="ACCEPT-CASE-08",
            client_name="上海星澜智能科技有限公司",
            counterparty_name="深圳远川供应链有限公司",
            case_type="买卖合同纠纷",
            path=workspace_db,
        )
        batch_files = [
            BatchInputFile(path=path, file_name=path.name, file_size=path.stat().st_size)
            for path in sorted(case08.glob("batch_document*.docx"))
        ]
        batch_results = run_batch_review(
            batch_files,
            case08 / "base_info.xlsx",
            case08 / "evidence_index.xlsx",
            case_id=case.case_id,
            db_path=workspace_db,
        )
        successes = [result for result in batch_results if result.status == "成功"]
        if len(successes) < 2:
            raise SystemExit(f"Case 08 批量复核成功数量不足：{len(successes)}")
        differences = build_case_field_differences(case.case_id, str(workspace_db))
        if not differences:
            raise SystemExit("Case 08 未生成案件内字段差异提示")
        case_package = build_case_review_package_zip(case.case_id, str(workspace_db))
        with ZipFile(BytesIO(case_package)) as package:
            package_names = set(package.namelist())
        if "案件复核汇总.xlsx" not in package_names or "案件复核报告.html" not in package_names:
            raise SystemExit("案件级复核包缺少汇总 Excel 或 HTML 报告")
        history_rows = list_review_runs(case.case_id, workspace_db)
        if len(history_rows) < len(successes):
            raise SystemExit("案件历史记录数量少于成功复核数量")

    with TemporaryDirectory() as tmpdir:
        package_dir = build_portable(Path(tmpdir))
        if not (package_dir / "start.bat").exists():
            raise SystemExit("portable 包缺少 start.bat")
        if not (package_dir / "README_快速启动.md").exists():
            raise SystemExit("portable 包缺少 README_快速启动.md")
        portable_paths = [path.relative_to(package_dir).as_posix() for path in package_dir.rglob("*")]
        if any(
            path == ".git"
            or path.startswith(".git/")
            or path == ".venv"
            or path.startswith(".venv/")
            or "__pycache__" in path
            or path.endswith("review_runs.sqlite")
            for path in portable_paths
        ):
            raise SystemExit("portable 包包含用户数据或缓存文件")

    with TemporaryDirectory() as tmpdir:
        preflight_root = Path(tmpdir) / "preflight"
        (preflight_root / "app").mkdir(parents=True)
        (preflight_root / "app" / "main.py").write_text("print('ok')", encoding="utf-8")
        (preflight_root / "requirements.txt").write_text("", encoding="utf-8")
        run_preflight(preflight_root)
        if not (preflight_root / "logs" / "current_port.txt").exists():
            raise SystemExit("preflight check 未写入 current_port.txt")

    with TemporaryDirectory() as tmpdir:
        zip_path = build_release_zip(version="acceptance", dist_root=Path(tmpdir))
        if not zip_path.exists():
            raise SystemExit("release zip 未生成")
        with ZipFile(zip_path) as release_zip:
            release_names = set(release_zip.namelist())
        required_release_files = {
            "LegalFactReview/start.bat",
            "LegalFactReview/README_快速启动.md",
            "LegalFactReview/app/config/review_rules.yml",
            "LegalFactReview/docs/规则配置说明.md",
            "LegalFactReview/docs/用户使用说明书.md",
        }
        if not required_release_files.issubset(release_names):
            raise SystemExit("release zip 缺少启动文件或快速说明")
        if not any(name.startswith("LegalFactReview/app/sample_data/test_cases/") for name in release_names):
            raise SystemExit("release zip 缺少 sample test cases")
        if not any(name.startswith("LegalFactReview/app/") for name in release_names):
            raise SystemExit("release zip 缺少 app 目录")
        if any(
            ".venv" in name
            or ".git/" in name
            or "__pycache__" in name
            or name.endswith("review_runs.sqlite")
            for name in release_names
        ):
            raise SystemExit("release zip 包含用户数据或缓存文件")

    with TemporaryDirectory() as tmpdir:
        blank_pdf = Path(tmpdir) / "blank_scan.pdf"
        doc = fitz.open()
        doc.new_page(width=200, height=200)
        doc.save(blank_pdf)
        doc.close()
        if not is_scanned_pdf(blank_pdf):
            raise SystemExit("PDF 扫描件检测未识别空白扫描型 PDF")
        try:
            parse_pdf_with_ocr(blank_pdf)
        except OcrDependencyError as exc:
            if "requirements-ocr.txt" not in str(exc):
                raise SystemExit("OCR 依赖缺失提示不够清晰")
        except Exception as exc:
            raise SystemExit(f"OCR 依赖缺失时未返回友好提示：{exc}") from exc

    print(
        "Acceptance check passed: "
        f"full={len(findings)} findings/{len(risk_types)} risk types; "
        f"single={len(single_findings)} findings/{len(external_fields)} external fields."
    )


def _extract_docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return "\n".join(texts)


if __name__ == "__main__":
    main()

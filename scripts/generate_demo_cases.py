from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "demo_cases" / "files"
DEMO_CASE_VERSION = "demo-cases-v3"
VERSION_FILE = OUTPUT_DIR / ".version"


def generate_demo_case_files(output_dir: str | Path = OUTPUT_DIR) -> list[Path]:
    root = Path(output_dir)
    root.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []

    generated.extend(
        _write_case(
            root / "clean_contract_pass",
            "clean_contract_pass.docx",
            [
                "民事起诉状",
                "案号：（2026）沪0115民初60001号",
                "原告：上海星河科技股份有限公司，统一社会信用代码：91310115MA1RXH0026。",
                "法定代表人：沈牧舟。住所地：上海市浦东新区星港路88号12层。",
                "被告：陈思齐。",
                "地址：上海市徐汇区青桥路12号。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2024-SH-001。",
                "2026年1月10日，双方签署《技术服务合同》。",
                "合同约定付款期限及履行截止日期为2026年2月10日。",
                "诉讼请求：请求判令被告支付服务费人民币280,000元及利息，利息自2026年3月1日起算。",
                "事实与理由：被告尚欠服务费人民币280,000元，原告于2026年4月1日起诉。",
                "证据1：技术服务合同；证据2：验收单；证据3：付款记录。",
                "标准条款：双方应按约定完成付款与结算；甲方收到合法有效发票后付款。",
                "标准条款：双方对合作信息承担保密义务。",
                "标准条款：验收完成后乙方开具发票，甲方按流程确认。",
            ],
            base_rows=[
                ("当事人名称", "上海星河科技股份有限公司", "2026-06-01", "synthetic"),
                ("法定代表人", "沈牧舟", "2026-06-01", "synthetic"),
                ("地址", "上海市浦东新区星港路88号12层", "2026-06-01", "synthetic"),
                ("统一社会信用代码", "91310115MA1RXH0026", "2026-06-01", "synthetic"),
                ("案号", "（2026）沪0115民初60001号", "2026-06-01", "synthetic"),
            ],
            evidence_rows=[
                ("1", "技术服务合同", "附件1", "synthetic"),
                ("2", "验收单", "附件2", "synthetic"),
                ("3", "付款记录", "附件3", "synthetic"),
            ],
        )
    )

    generated.extend(
        _write_case(
            root / "missing_required_fields",
            "missing_required_fields.docx",
            [
                "合同审查记录",
                "甲方：上海星河科技股份有限公司，统一社会信用代码：91310115MA1RXH0026。",
                "乙方：杭州青岚贸易有限公司。",
                "签署日期：2026年2月10日。",
                "合同金额：人民币32O,OOO元。",
                "本记录故意缺少合同编号、法定代表人和注册地址。",
            ],
            base_rows=[
                ("当事人名称", "上海星河科技股份有限公司", "2026-06-01", "synthetic"),
                ("法定代表人", "沈牧舟", "2026-06-01", "synthetic"),
                ("地址", "上海市浦东新区星港路88号12层", "2026-06-01", "synthetic"),
                ("统一社会信用代码", "91310115MA1RXH0026", "2026-06-01", "synthetic"),
                ("合同编号", "HT-2024-SH-001", "2026-06-01", "synthetic"),
            ],
            evidence_rows=[
                ("1", "合同审查底稿", "附件1", "synthetic"),
            ],
        )
    )

    generated.extend(
        _write_case(
            root / "client_profile_mismatch",
            "client_profile_mismatch.docx",
            [
                "民事起诉状",
                "案号：（2026）沪0115民初60002号",
                "原告：上海澄川贸易有限公司，统一社会信用代码：91310115MA1RXH0026。",
                "法定代表人：刘远航。住所地：上海市浦东新区澄川路18号。",
                "被告：杭州青岚贸易有限公司。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2024-SH-002。",
                "2026年2月1日，双方签署《技术服务合同》。",
                "诉讼请求：请求判令被告支付人民币180,000元。",
                "事实与理由：证据1为合同，证据2为验收单。",
            ],
            base_rows=[
                ("当事人名称", "上海澄川贸易有限公司", "2026-06-01", "synthetic"),
                ("法定代表人", "刘远航", "2026-06-01", "synthetic"),
                ("地址", "上海市浦东新区澄川路18号", "2026-06-01", "synthetic"),
                ("统一社会信用代码", "91310115MA1RXH0026", "2026-06-01", "synthetic"),
                ("案号", "（2026）沪0115民初60002号", "2026-06-01", "synthetic"),
            ],
            evidence_rows=[
                ("1", "技术服务合同", "附件1", "synthetic"),
                ("2", "验收单", "附件2", "synthetic"),
            ],
        )
    )

    generated.extend(
        _write_case(
            root / "missing_standard_clause",
            "missing_standard_clause.docx",
            [
                "合同审查记录",
                "甲方：上海星河科技股份有限公司，统一社会信用代码：91310115MA1RXH0026。",
                "乙方：深圳远川供应链有限公司。",
                "合同编号：HT-2024-SH-003。",
                "合同金额：人民币430,000元。",
                "签署日期：2026年1月8日。",
                "服务内容：乙方向甲方提供企业软件实施服务。",
                "付款安排：甲方在验收后支付服务费。",
                "本合同仅保留核心商务安排，故意省略若干平台治理条款以供评估。",
            ],
            base_rows=None,
            evidence_rows=None,
        )
    )

    generated.extend(
        _write_case(
            root / "llm_fallback_privacy_export",
            "llm_fallback_privacy_export.docx",
            [
                "律师函",
                "致：杭州青岚贸易有限公司",
                "委托人：上海星河科技股份有限公司，统一社会信用代码：91310115MA1RXH0026。",
                "法定代表人：沈牧舟。地址：上海市浦东新区星港路88号12层。",
                "请贵方就合同编号 HT-2024-SH-004 项下人民币88,000元款项于2026年6月30日前反馈。",
                "联系邮箱：demo-contact@example.com；联系电话：13812345678。",
                "本函仅为 synthetic demo，不构成法律意见。",
            ],
            base_rows=None,
            evidence_rows=None,
        )
    )
    VERSION_FILE.write_text(DEMO_CASE_VERSION, encoding="utf-8")
    generated.append(VERSION_FILE)
    return generated


def _write_case(
    case_dir: Path,
    document_name: str,
    lines: list[str],
    base_rows: list[tuple[str, str, str, str]] | None,
    evidence_rows: list[tuple[str, str, str, str]] | None,
) -> list[Path]:
    case_dir.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []
    document_path = case_dir / document_name
    _write_docx(document_path, lines)
    generated.append(document_path)
    if base_rows is not None:
        path = case_dir / "base_info.xlsx"
        _write_workbook(path, ("字段名", "标准值", "更新时间", "备注"), base_rows)
        generated.append(path)
    if evidence_rows is not None:
        path = case_dir / "evidence_index.xlsx"
        _write_workbook(path, ("证据编号", "证据名称", "页码/附件位置", "备注"), evidence_rows)
        generated.append(path)
    return generated


def _write_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for index, line in enumerate(lines):
        if index == 0:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    tmp_path = path.with_name(f"{path.stem}.__tmp__{path.suffix}")
    document.save(tmp_path)
    tmp_path.replace(path)


def _write_workbook(path: Path, headers: tuple[str, str, str, str], rows: list[tuple[str, str, str, str]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "synthetic"
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(list(row))
    tmp_path = path.with_name(f"{path.stem}.__tmp__{path.suffix}")
    workbook.save(tmp_path)
    tmp_path.replace(path)


def main() -> None:
    generated = generate_demo_case_files()
    print(f"Generated {len(generated)} synthetic demo case files in {OUTPUT_DIR}")
    for path in generated:
        print(path.relative_to(ROOT).as_posix())


if __name__ == "__main__":
    main()

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_DIR = ROOT / "app" / "sample_data"


def make_document() -> None:
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("民事起诉状", level=1)
    doc.add_paragraph("案号：（202X）沪0101民初ABC号")
    doc.add_paragraph("原告：上海星河科技有限公司，统一社会信用代码：91310000MA1K000001。")
    doc.add_paragraph("法定代表人：李四。住所地：上海市浦东新区示例路88号。")
    doc.add_paragraph("被告：杭州蓝海贸易有限公司，统一社会信用代码：91330100MA2K000002。")
    doc.add_paragraph("合同编号：HT-2024-001。")
    doc.add_paragraph("2024年5月10日，原告上海星河科技股份有限公司与被告签署《采购合同》。")
    doc.add_paragraph("合同约定付款期限及履行截止日期为2024年4月30日。")
    doc.add_paragraph("诉讼请求：请求判令被告支付货款人民币100,000元及利息，利息自2024年6月1日起算。")
    doc.add_paragraph("事实与理由：被告尚欠货款人民币120,000元，原告于2024年5月1日起诉。")
    doc.add_paragraph("证据1：采购合同；证据2：发票；证据7：付款催告函。")
    doc.add_paragraph("基础信息版本日期：2024年1月1日。")
    doc.add_paragraph("其他事项：待补充。")
    doc.save(SAMPLE_DIR / "sample_document.docx")


def make_base_info() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "基础信息"
    ws.append(["字段名", "标准值", "更新时间", "备注"])
    ws.append(["当事人名称", "上海星河科技有限公司", "2024-06-15", "虚构公司"])
    ws.append(["法定代表人", "张三", "2024-06-15", "文书中故意写成李四"])
    ws.append(["地址", "上海市浦东新区示例路99号", "2024-06-15", "文书中地址不一致"])
    ws.append(["统一社会信用代码", "91310000MA1K000001", "2024-06-15", "虚构代码"])
    ws.append(["案号", "（2024）沪0101民初1234号", "2024-06-15", "文书中案号异常"])
    wb.save(SAMPLE_DIR / "sample_base_info.xlsx")


def make_evidence_index() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "证据目录"
    ws.append(["证据编号", "证据名称", "页码/附件位置", "备注"])
    ws.append(["1", "采购合同", "附件1", "虚构材料"])
    ws.append(["2", "发票", "附件2", "虚构材料"])
    ws.append(["3", "往来邮件", "附件3", "虚构材料"])
    wb.save(SAMPLE_DIR / "sample_evidence_index.xlsx")


if __name__ == "__main__":
    make_document()
    make_base_info()
    make_evidence_index()
    print(f"Sample data generated in {SAMPLE_DIR}")

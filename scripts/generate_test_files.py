from __future__ import annotations

import shutil
import textwrap
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
TEST_CASES_DIR = ROOT / "app" / "sample_data" / "test_cases"


@dataclass(frozen=True)
class WorkbookRow:
    values: tuple[str, str, str, str]


@dataclass(frozen=True)
class CaseSpec:
    dirname: str
    readme: str
    documents: dict[str, list[str]]
    base_info: list[WorkbookRow] | None = None
    evidence_index: list[WorkbookRow] | None = None


def _paragraphs(*lines: str) -> list[str]:
    return [line.strip() for line in lines if line.strip()]


CASE_SPECS: list[CaseSpec] = [
    CaseSpec(
        dirname="case_01_full_review_basic",
        documents={
            "sample_complaint.docx": _paragraphs(
                "民事起诉状",
                "案号：（202X）沪0115民初ABC号",
                "原告：上海星澜智能科技有限公司，统一社会信用代码：91310000MA1K000001。",
                "法定代表人：张明远。住所地：上海市浦东新区星澜路88号。",
                "被告：北京云衡数据服务有限公司，统一社会信用代码：91310000MA1K000002。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2026-STAR-001。",
                "2026年5月20日，原告上海星澜智慧科技有限公司与被告签署《数据服务合同》。",
                "合同约定付款期限及履行截止日期为2026年5月1日。",
                "诉讼请求：请求判令被告支付服务费人民币1,280,000元及利息，利息自2026年6月1日起算。",
                "事实与理由：被告尚欠服务费人民币1,208,000元，原告于2026年5月10日起诉。",
                "证据1：数据服务合同；证据2：验收单；证据7：付款催告函。",
                "基础信息版本日期：2026年1月1日。",
                "其他事项：待补充。",
            ),
            "sample_complaint.pdf": _paragraphs(
                "民事起诉状",
                "案号：（202X）沪0115民初ABC号",
                "原告：上海星澜智能科技有限公司，统一社会信用代码：91310000MA1K000001。",
                "法定代表人：张明远。住所地：上海市浦东新区星澜路88号。",
                "被告：北京云衡数据服务有限公司，统一社会信用代码：91310000MA1K000002。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2026-STAR-001。",
                "2026年5月20日，原告上海星澜智慧科技有限公司与被告签署《数据服务合同》。",
                "合同约定付款期限及履行截止日期为2026年5月1日。",
                "诉讼请求：请求判令被告支付服务费人民币1,280,000元及利息，利息自2026年6月1日起算。",
                "事实与理由：被告尚欠服务费人民币1,208,000元，原告于2026年5月10日起诉。",
                "证据1：数据服务合同；证据2：验收单；证据7：付款催告函。",
                "基础信息版本日期：2026年1月1日。",
                "其他事项：待补充。",
            ),
        },
        base_info=[
            WorkbookRow(("当事人名称", "上海星澜智慧科技有限公司", "2026-06-01", "虚构标准主体")),
            WorkbookRow(("法定代表人", "李承安", "2026-06-01", "文书中故意不一致")),
            WorkbookRow(("地址", "上海市浦东新区星澜路99号", "2026-06-01", "文书中地址不一致")),
            WorkbookRow(("统一社会信用代码", "91310000MA1K000001", "2026-06-01", "虚构代码")),
            WorkbookRow(("案号", "（2026）沪0115民初12345号", "2026-06-01", "文书中案号异常")),
        ],
        evidence_index=[
            WorkbookRow(("1", "数据服务合同", "附件1", "虚构材料")),
            WorkbookRow(("2", "验收单", "附件2", "虚构材料")),
            WorkbookRow(("3", "发票", "附件3", "虚构材料")),
            WorkbookRow(("4", "对账单", "附件4", "虚构材料")),
            WorkbookRow(("5", "往来邮件", "附件5", "虚构材料")),
        ],
        readme="""
        # Case 01：完整复核基础场景

        ## 上传方式

        上传 `sample_complaint.docx` 或 `sample_complaint.pdf`，并同时上传 `base_info.xlsx` 和 `evidence_index.xlsx`。

        ## 预期模式

        完整复核模式。

        ## 预期结果

        - 生成多个 findings。
        - 至少覆盖 SOURCE_MISMATCH、AMOUNT_INCONSISTENCY、EVIDENCE_REFERENCE_MISSING、PLACEHOLDER_REMAINING、FORMAT_ERROR、DATE_LOGIC_WARNING。
        - 可用于演示基础信息比对、证据缺失、金额不一致和日期顺序可疑。

        ## 不应出现

        - 不应因 PDF 是扫描件而要求 OCR。
        - 不应包含真实公司、真实案号或真实个人信息。
        """,
    ),
    CaseSpec(
        dirname="case_02_single_document_only",
        documents={
            "single_document.docx": _paragraphs(
                "民事起诉状",
                "案号：（2026）京0105民初23456号",
                "原告：北京云衡数据服务有限公司，统一社会信用代码：91310000MA1K000002。",
                "法定代表人：王若宁。住所地：北京市朝阳区云衡路66号。",
                "被告：杭州青岚贸易有限公司，统一社会信用代码：91310000MA1K000003。",
                "地址：杭州市滨江区青岚街18号。",
                "管辖法院：北京市朝阳区人民法院。",
                "合同编号：HT-2026-CLOUD-002。",
                "2026年3月15日，双方签署《云数据服务合同》。",
                "诉讼请求：请求判令被告支付人民币560,000元。",
                "事实与理由：证据1为合同，证据2为验收单，证据3为付款记录。",
            ),
            "single_document.pdf": _paragraphs(
                "民事起诉状",
                "案号：（2026）京0105民初23456号",
                "原告：北京云衡数据服务有限公司，统一社会信用代码：91310000MA1K000002。",
                "法定代表人：王若宁。住所地：北京市朝阳区云衡路66号。",
                "被告：杭州青岚贸易有限公司，统一社会信用代码：91310000MA1K000003。",
                "地址：杭州市滨江区青岚街18号。",
                "管辖法院：北京市朝阳区人民法院。",
                "合同编号：HT-2026-CLOUD-002。",
                "2026年3月15日，双方签署《云数据服务合同》。",
                "诉讼请求：请求判令被告支付人民币560,000元。",
                "事实与理由：证据1为合同，证据2为验收单，证据3为付款记录。",
            ),
        },
        readme="""
        # Case 02：单文书字段抽取模式

        ## 上传方式

        只上传 `single_document.docx` 或 `single_document.pdf`，不要上传基础信息表和证据目录。

        ## 预期模式

        单文书字段抽取模式。

        ## 预期结果

        - 生成待外部复核字段清单。
        - 不生成 SOURCE_MISMATCH。
        - 不因缺少证据目录生成 EVIDENCE_REFERENCE_MISSING。
        - 可用于演示主体、案号、金额、日期、合同编号和证据编号抽取。

        ## 不应出现

        - 不应提示必须上传基础信息表或证据目录。
        """,
    ),
    CaseSpec(
        dirname="case_03_many_conflicts",
        documents={
            "conflict_document.docx": _paragraphs(
                "民事起诉状",
                "案号：（2026）浙0108民初34567号",
                "案号：（2026）浙0108民初34576号",
                "原告：上海星澜智能科技有限公司，统一社会信用代码：91310000MA1K000001。",
                "原告：上海星澜智慧科技有限公司。",
                "甲方：上海星澜智能技术有限公司。",
                "法定代表人：张明远。",
                "法定代表人：李承安。",
                "被告：深圳远川供应链有限公司，统一社会信用代码：91310000MA1K000004。",
                "合同编号：HT-2026-STAR-001。",
                "合同编号：HT-2026-STAR-009。",
                "2026年4月2日，双方签署合作协议。",
                "诉讼请求：请求支付人民币1,280,000元。",
                "事实与理由：欠款金额为人民币1,208,000元，补充表述为人民币128,000元。",
                "证据1：合同；证据3：验收单；证据6：付款记录；证据9：催告函。",
                "附件1为合同复印件，附件4为某某确认单。",
                "其他事项：XXXX；待补充；签章栏：【 】。",
            ),
        },
        base_info=[
            WorkbookRow(("当事人名称", "上海星澜智能科技有限公司", "2026-06-01", "虚构标准主体")),
            WorkbookRow(("法定代表人", "张明远", "2026-06-01", "虚构人员")),
            WorkbookRow(("统一社会信用代码", "91310000MA1K000001", "2026-06-01", "虚构代码")),
            WorkbookRow(("案号", "（2026）浙0108民初34567号", "2026-06-01", "虚构案号")),
            WorkbookRow(("地址", "上海市浦东新区星澜路88号", "2026-06-01", "补充标准地址")),
        ],
        evidence_index=[
            WorkbookRow(("1", "合同", "附件1", "虚构材料")),
            WorkbookRow(("2", "发票", "附件2", "虚构材料")),
            WorkbookRow(("3", "验收单", "附件3", "虚构材料")),
        ],
        readme="""
        # Case 03：高冲突文书

        ## 上传方式

        上传 `conflict_document.docx`，并同时上传 `base_info.xlsx` 和 `evidence_index.xlsx`。

        ## 预期模式

        完整复核模式。

        ## 预期结果

        - INTERNAL_CONFLICT 明显。
        - AMOUNT_INCONSISTENCY 明显。
        - PLACEHOLDER_REMAINING 明显。
        - high 风险数量较多。

        ## 不应出现

        - 不应因证据目录跳号或重复导致程序崩溃。
        """,
    ),
    CaseSpec(
        dirname="case_04_missing_fields",
        documents={
            "missing_fields_document.docx": _paragraphs(
                "民事起诉状",
                "原告上海星澜智能科技有限公司与被告杭州青岚贸易有限公司存在买卖合同争议。",
                "住所地：上海市浦东新区星澜路88号。",
                "2026年2月10日，双方签署供货安排并完成部分交付。",
                "诉讼请求：请求判令被告支付人民币320,000元。",
                "事实与理由：证据1为送货单，证据2为对账单。",
            ),
        },
        base_info=[
            WorkbookRow(("当事人名称", "上海星澜智能科技有限公司", "2026-06-01", "文书未用标签列明")),
            WorkbookRow(("法定代表人", "张明远", "2026-06-01", "文书缺失")),
            WorkbookRow(("地址", "上海市浦东新区星澜路88号", "2026-06-01", "文书有地址")),
            WorkbookRow(("统一社会信用代码", "91310000MA1K000001", "2026-06-01", "文书缺失")),
            WorkbookRow(("案号", "（2026）沪0115民初12345号", "2026-06-01", "文书缺失")),
        ],
        evidence_index=[
            WorkbookRow(("1", "送货单", "附件1", "虚构材料")),
            WorkbookRow(("2", "对账单", "附件2", "虚构材料")),
        ],
        readme="""
        # Case 04：关键字段缺失

        ## 上传方式

        上传 `missing_fields_document.docx`，并同时上传 `base_info.xlsx` 和 `evidence_index.xlsx`。

        ## 预期模式

        完整复核模式。

        ## 预期结果

        - 生成 MISSING_FIELD。
        - 仍能抽取已有金额、日期、地址和证据编号。
        - 报告导出不应崩溃。
        - source_location 不得为空。

        ## 不应出现

        - 不应因为缺少案号、法院名称或合同编号导致解析失败。
        """,
    ),
    CaseSpec(
        dirname="case_05_evidence_missing",
        documents={
            "evidence_problem_document.docx": _paragraphs(
                "民事起诉状",
                "案号：（2026）沪0115民初22334号",
                "原告：深圳远川供应链有限公司，统一社会信用代码：91310000MA1K000004。",
                "法定代表人：赵嘉禾。住所地：深圳市南山区远川路20号。",
                "被告：杭州青岚贸易有限公司，统一社会信用代码：91310000MA1K000003。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2026-SUPPLY-003。",
                "2026年1月8日，双方签署供应链服务合同。",
                "诉讼请求：请求判令被告支付人民币430,000元。",
                "事实与理由：证据1为合同，证据2为对账单，证据4为签收单，证据8为催告函。",
                "附件1为合同副本，附件8为催告函扫描件。",
            ),
        },
        base_info=[
            WorkbookRow(("当事人名称", "深圳远川供应链有限公司", "2026-06-01", "虚构标准主体")),
            WorkbookRow(("法定代表人", "赵嘉禾", "2026-06-01", "虚构人员")),
            WorkbookRow(("统一社会信用代码", "91310000MA1K000004", "2026-06-01", "虚构代码")),
            WorkbookRow(("案号", "（2026）沪0115民初22334号", "2026-06-01", "虚构案号")),
            WorkbookRow(("地址", "深圳市南山区远川路20号", "2026-06-01", "虚构地址")),
        ],
        evidence_index=[
            WorkbookRow(("1", "供应链服务合同", "附件1", "虚构材料")),
            WorkbookRow(("2", "对账单", "附件2", "虚构材料")),
            WorkbookRow(("2", "重复编号测试行", "附件2-重复", "故意重复")),
            WorkbookRow(("3", "发票", "附件3", "虚构材料")),
        ],
        readme="""
        # Case 05：证据目录问题

        ## 上传方式

        上传 `evidence_problem_document.docx`，并同时上传 `base_info.xlsx` 和 `evidence_index.xlsx`。

        ## 预期模式

        完整复核模式。

        ## 预期结果

        - 生成 EVIDENCE_REFERENCE_MISSING。
        - 能展示缺失的证据4、证据8或附件8。
        - 证据目录存在重复编号时程序不应崩溃。

        ## 不应出现

        - 不应把证据目录重复行当作致命解析错误。
        """,
    ),
    CaseSpec(
        dirname="case_06_clean_document",
        documents={
            "clean_document.docx": _paragraphs(
                "民事起诉状",
                "案号：（2026）沪0115民初12345号",
                "原告：上海星澜智能科技有限公司，统一社会信用代码：91310000MA1K000001。",
                "法定代表人：张明远。住所地：上海市浦东新区星澜路88号。",
                "被告：陈思齐。",
                "地址：上海市徐汇区青桥路12号。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2026-STAR-001。",
                "2026年1月10日，原告与被告签署《技术服务合同》。",
                "合同约定付款期限及履行截止日期为2026年2月10日。",
                "诉讼请求：请求判令被告支付服务费人民币280,000元及利息，利息自2026年3月1日起算。",
                "事实与理由：被告尚欠服务费人民币280,000元，原告于2026年4月1日起诉。",
                "证据1：技术服务合同；证据2：验收单；证据3：付款记录。",
            ),
        },
        base_info=[
            WorkbookRow(("当事人名称", "上海星澜智能科技有限公司", "2026-06-01", "虚构标准主体")),
            WorkbookRow(("法定代表人", "张明远", "2026-06-01", "虚构人员")),
            WorkbookRow(("地址", "上海市浦东新区星澜路88号", "2026-06-01", "虚构地址")),
            WorkbookRow(("统一社会信用代码", "91310000MA1K000001", "2026-06-01", "虚构代码")),
            WorkbookRow(("案号", "（2026）沪0115民初12345号", "2026-06-01", "虚构案号")),
        ],
        evidence_index=[
            WorkbookRow(("1", "技术服务合同", "附件1", "虚构材料")),
            WorkbookRow(("2", "验收单", "附件2", "虚构材料")),
            WorkbookRow(("3", "付款记录", "附件3", "虚构材料")),
        ],
        readme="""
        # Case 06：基本干净文书

        ## 上传方式

        上传 `clean_document.docx`，并同时上传 `base_info.xlsx` 和 `evidence_index.xlsx`。

        ## 预期模式

        完整复核模式。

        ## 预期结果

        - findings 数量较少。
        - 可用于测试“无明显风险”或“低风险”展示。
        - 系统应正常导出 Excel、HTML、Word 和 ZIP。

        ## 不应出现

        - 不应出现 SOURCE_MISMATCH。
        - 不应出现 EVIDENCE_REFERENCE_MISSING。
        - 不应出现 PLACEHOLDER_REMAINING。
        """,
    ),
    CaseSpec(
        dirname="case_07_pdf_text_extraction",
        documents={
            "text_pdf_document.pdf": _paragraphs(
                "PDF 文本抽取测试材料",
                "案号：（2026）京0105民初23456号",
                "原告：北京云衡数据服务有限公司，统一社会信用代码：91310000MA1K000002。",
                "法定代表人：王若宁。住所地：北京市朝阳区云衡路66号。",
                "被告：杭州青岚贸易有限公司。",
                "合同编号：HT-2026-CLOUD-002。",
                "证据1：云服务合同；证据2：验收单。",
            ),
        },
        readme="""
        # Case 07：PDF 文本抽取测试

        ## 上传方式

        只上传 `text_pdf_document.pdf`。

        ## 预期模式

        单文书字段抽取模式。

        ## 预期结果

        - PDF 中应有可复制文本。
        - 可用于验证普通 PDF 文本解析。
        - 不需要 OCR。

        ## 不应出现

        - 不应被识别为空白扫描件。
        """,
    ),
    CaseSpec(
        dirname="case_08_case_workspace_batch",
        documents={
            "batch_document_v1.docx": _paragraphs(
                "民事起诉状",
                "案号：（2026）沪0115民初45678号",
                "原告：上海星澜智能科技有限公司，统一社会信用代码：91310000MA1K000001。",
                "法定代表人：张明远。住所地：上海市浦东新区星澜路88号。",
                "被告：深圳远川供应链有限公司，统一社会信用代码：91310000MA1K000004。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2026-STAR-001。",
                "2026年2月1日，双方签署《技术服务合同》。",
                "诉讼请求：请求判令被告支付人民币680,000元。",
                "事实与理由：证据1为合同，证据2为验收单。",
            ),
            "batch_document_v2.docx": _paragraphs(
                "民事起诉状修订稿",
                "案号：（2026）沪0115民初45678号",
                "原告：上海星澜智慧科技有限公司，统一社会信用代码：91310000MA1K000001。",
                "法定代表人：张明远。住所地：上海市浦东新区星澜路88号。",
                "被告：深圳远川供应链有限公司，统一社会信用代码：91310000MA1K000004。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2026-STAR-001。",
                "2026年2月1日，双方签署《技术服务合同》。",
                "诉讼请求：请求判令被告支付人民币860,000元。",
                "事实与理由：证据1为合同，证据3为补充验收单，证据8为催告函。",
                "其他事项：待补充。",
            ),
            "batch_document_clean.docx": _paragraphs(
                "民事起诉状定稿",
                "案号：（2026）沪0115民初45678号",
                "原告：上海星澜智能科技有限公司，统一社会信用代码：91310000MA1K000001。",
                "法定代表人：张明远。住所地：上海市浦东新区星澜路88号。",
                "被告：深圳远川供应链有限公司，统一社会信用代码：91310000MA1K000004。",
                "管辖法院：上海市浦东新区人民法院。",
                "合同编号：HT-2026-STAR-001。",
                "2026年2月1日，双方签署《技术服务合同》。",
                "合同约定付款期限及履行截止日期为2026年3月1日。",
                "诉讼请求：请求判令被告支付人民币680,000元。",
                "事实与理由：证据1为合同，证据2为验收单，证据3为付款记录。",
            ),
        },
        base_info=[
            WorkbookRow(("当事人名称", "上海星澜智能科技有限公司", "2026-06-01", "虚构标准主体")),
            WorkbookRow(("法定代表人", "张明远", "2026-06-01", "虚构人员")),
            WorkbookRow(("地址", "上海市浦东新区星澜路88号", "2026-06-01", "虚构地址")),
            WorkbookRow(("统一社会信用代码", "91310000MA1K000001", "2026-06-01", "虚构代码")),
            WorkbookRow(("案号", "（2026）沪0115民初45678号", "2026-06-01", "虚构案号")),
        ],
        evidence_index=[
            WorkbookRow(("1", "技术服务合同", "附件1", "虚构材料")),
            WorkbookRow(("2", "验收单", "附件2", "虚构材料")),
            WorkbookRow(("3", "付款记录", "附件3", "虚构材料")),
        ],
        readme="""
        # Case 08：案件工作区批量复核

        ## 上传方式

        先创建一个案件工作区，再在批量复核中同时上传 `batch_document_v1.docx`、`batch_document_v2.docx` 和 `batch_document_clean.docx`，并上传 `base_info.xlsx` 与 `evidence_index.xlsx`。

        ## 预期模式

        完整复核模式，且所有 review_id 关联到同一个案件。

        ## 预期结果

        - 至少 2 份文书复核成功。
        - `batch_document_v2.docx` 会出现主体名称、金额、证据编号和占位符相关提示。
        - 案件内字段差异提示应展示主体名称和金额在不同文书间不一致。
        - `batch_document_clean.docx` 可作为较干净文书对照。

        ## 不应出现

        - 单份文书失败不应影响其他文书生成 review_id。
        - 案件级复核包不应包含完整原文书内容。
        """,
    ),
    CaseSpec(
        dirname="case_09_custom_rules_profile",
        documents={
            "lawyer_letter.docx": _paragraphs(
                "律师函",
                "致：杭州青岚贸易有限公司",
                "委托人：上海星澜智能科技有限公司，统一社会信用代码：91310000MA1K000001。",
                "法定代表人：张明远。地址：上海市浦东新区星澜路88号。",
                "我方就合同编号：HT-2026-STAR-001 项下人民币680,000元款项事宜致函。",
                "请贵方于2026年6月30日前反馈。【请律师确认】",
            ),
            "contract_review.docx": _paragraphs(
                "合同审查记录",
                "甲方：上海星澜智能科技有限公司。",
                "乙方：深圳远川供应链有限公司。",
                "合同编号：HT-2026-SUPPLY-003。",
                "合同金额：人民币430,000元。",
                "签署日期：2026年1月8日。",
            ),
        },
        readme="""
        # Case 09：自定义规则与文书类型模板

        ## 上传方式

        在“规则配置中心”导入 `custom_review_rules.yml`，再上传 `lawyer_letter.docx` 或 `contract_review.docx` 运行复核。

        ## 预期模式

        可选择律师函或合同审查模板。

        ## 预期结果

        - 自定义配置禁用 PLACEHOLDER_REMAINING 后，占位符不再输出为 finding。
        - SOURCE_MISMATCH 默认等级调整为 medium。
        - 律师函模板新增 required_fields，可影响 MISSING_FIELD。
        - 恢复默认配置后行为恢复。
        """,
    ),
]


def generate_test_files(clean: bool = True) -> list[Path]:
    if clean and TEST_CASES_DIR.exists():
        shutil.rmtree(TEST_CASES_DIR)
    TEST_CASES_DIR.mkdir(parents=True, exist_ok=True)

    generated: list[Path] = []
    for spec in CASE_SPECS:
        case_dir = TEST_CASES_DIR / spec.dirname
        case_dir.mkdir(parents=True, exist_ok=True)

        for filename, lines in spec.documents.items():
            path = case_dir / filename
            if path.suffix.lower() == ".docx":
                _write_docx(path, lines)
            elif path.suffix.lower() == ".pdf":
                _write_pdf(path, lines)
            else:
                raise ValueError(f"Unsupported document type: {path}")
            generated.append(path)

        if spec.base_info is not None:
            path = case_dir / "base_info.xlsx"
            _write_workbook(path, ("字段名", "标准值", "更新时间", "备注"), spec.base_info)
            generated.append(path)

        if spec.evidence_index is not None:
            path = case_dir / "evidence_index.xlsx"
            _write_workbook(path, ("证据编号", "证据名称", "页码/附件位置", "备注"), spec.evidence_index)
            generated.append(path)

        readme_path = case_dir / "README.md"
        readme_path.write_text(textwrap.dedent(spec.readme).strip() + "\n", encoding="utf-8")
        generated.append(readme_path)

        if spec.dirname == "case_09_custom_rules_profile":
            config_path = case_dir / "custom_review_rules.yml"
            config_path.write_text(_case_09_custom_config(), encoding="utf-8")
            generated.append(config_path)

    return generated


def _write_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for index, line in enumerate(lines):
        if index == 0:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    document.save(path)


def _write_pdf(path: Path, lines: list[str]) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    y = 72
    for line in lines:
        if y > 780:
            page = document.new_page(width=595, height=842)
            y = 72
        page.insert_text((72, y), line, fontname="china-s", fontsize=11)
        y += 24
    document.save(path)
    document.close()


def _write_workbook(path: Path, headers: tuple[str, str, str, str], rows: list[WorkbookRow]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "测试数据"
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(list(row.values))
    workbook.save(path)


def _case_09_custom_config() -> str:
    return textwrap.dedent(
        """
        config_version: case09-custom-v1
        updated_by: local_user
        enabled_risk_types:
          - MISSING_FIELD
          - INTERNAL_CONFLICT
          - SOURCE_MISMATCH
          - FORMAT_ERROR
          - POSSIBLY_OUTDATED
          - EVIDENCE_REFERENCE_MISSING
          - AMOUNT_INCONSISTENCY
          - DATE_LOGIC_WARNING
        critical_fields:
          - 当事人名称
          - 合同编号
          - 金额
        placeholder_patterns:
          - 待补充
          - 【请律师确认】
        severity_mapping:
          SOURCE_MISMATCH: medium
          AMOUNT_INCONSISTENCY: high
          EVIDENCE_REFERENCE_MISSING: high
        external_review_sources:
          company_name:
            sources: [企查查, 国家企业信用信息公示系统]
            suggested_action: 建议人工复核主体名称是否与客户材料一致。
          amount:
            sources: [合同, 对账单]
            suggested_action: 建议人工复核金额是否与合同和对账材料一致。
          contract_number:
            sources: [合同原件, 合同台账]
            suggested_action: 建议人工复核合同编号是否与合同原件一致。
          default:
            sources: [基础材料]
            suggested_action: 建议人工复核该字段是否与基础材料一致。
        document_profiles:
          general:
            display_name: 通用文书
            required_fields: [party_name, case_number]
          lawyer_letter:
            display_name: 律师函
            required_fields: [party_name, amount, contract_number]
          contract_review:
            display_name: 合同审查
            required_fields: [party_name, contract_number, amount, date]
        field_dictionary:
          party_name:
            display_name: 主体名称
            aliases: [委托人, 收函人, 甲方, 乙方]
            regex_patterns: []
            suggested_review_sources: [客户材料]
            default_severity: high
            required_by_default: true
          contract_number:
            display_name: 合同编号
            aliases: [合同编号]
            regex_patterns: []
            suggested_review_sources: [合同原件]
            default_severity: medium
            required_by_default: true
          amount:
            display_name: 金额
            aliases: [款项, 金额]
            regex_patterns: ['(?:人民币)?[0-9][0-9,]*(?:\\.\\d+)?\\s*元']
            suggested_review_sources: [合同, 对账单]
            default_severity: high
            required_by_default: true
        """
    ).strip() + "\n"


def main() -> None:
    generated = generate_test_files(clean=True)
    print(f"Generated {len(generated)} files in {TEST_CASES_DIR}")
    for path in generated:
        print(path.relative_to(ROOT).as_posix())


if __name__ == "__main__":
    main()

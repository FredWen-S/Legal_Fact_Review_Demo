from __future__ import annotations

import sqlite3
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from app.models.schemas import CaseWorkspace
from app.reports.case_package import build_case_review_package_zip, build_case_summary_excel
from app.storage.sqlite_store import (
    init_db,
    create_case_workspace,
    delete_case_workspace,
    get_case_workspace,
    get_review_run_detail,
    list_case_workspaces,
    list_review_runs,
    save_review_run,
)
from app.workflows.batch_review import BatchInputFile, run_batch_review
from app.workflows.case_differences import build_case_field_differences
from scripts.generate_test_files import TEST_CASES_DIR, generate_test_files


def test_case_workspace_schema(tmp_path: Path):
    case = create_case_workspace("示例案件", internal_case_no="CASE-001", client_name="上海星澜智能科技有限公司", path=tmp_path / "reviews.sqlite")
    assert isinstance(case, CaseWorkspace)
    assert case.case_name == "示例案件"
    assert case.internal_case_no == "CASE-001"


def test_sqlite_create_query_delete_case(tmp_path: Path):
    db_path = tmp_path / "reviews.sqlite"
    case = create_case_workspace("测试案件", internal_case_no="CASE-001", path=db_path)

    assert get_case_workspace(case.case_id, db_path).case_name == "测试案件"
    assert list_case_workspaces(db_path)

    delete_case_workspace(case.case_id, db_path)
    assert get_case_workspace(case.case_id, db_path) is None


def test_empty_db_and_repeated_init_are_idempotent(tmp_path: Path):
    db_path = tmp_path / "reviews.sqlite"

    init_db(db_path).close()
    init_db(db_path).close()

    conn = sqlite3.connect(db_path)
    try:
        tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    finally:
        conn.close()
    assert {"case_workspaces", "case_documents", "review_runs", "manual_review_records"}.issubset(tables)


def test_v01_database_migrates_and_old_save_review_call_still_works(tmp_path: Path):
    db_path = tmp_path / "old.sqlite"
    conn = sqlite3.connect(db_path)
    try:
        conn.execute(
            """
            CREATE TABLE review_runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                finding_count INTEGER NOT NULL,
                findings_json TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE manual_review_records (
                record_id TEXT PRIMARY KEY,
                review_id TEXT NOT NULL,
                field_id TEXT NOT NULL,
                field_type TEXT NOT NULL,
                extracted_value TEXT NOT NULL,
                confirmed_value TEXT,
                review_status TEXT NOT NULL,
                review_source TEXT,
                reviewer_note TEXT,
                source_location TEXT NOT NULL,
                reviewed_at TEXT
            )
            """
        )
        conn.commit()
    finally:
        conn.close()

    init_db(db_path).close()
    review_id = save_review_run([], path=db_path)
    rows = list_review_runs(path=db_path)

    assert review_id == rows[0].review_id
    assert rows[0].case_id is None


def test_delete_case_does_not_delete_unrelated_temporary_reviews(tmp_path: Path):
    db_path = tmp_path / "reviews.sqlite"
    case = create_case_workspace("待删除案件", path=db_path)
    case_review_id = save_review_run([], path=db_path, case_id=case.case_id, file_name="case.docx")
    temp_review_id = save_review_run([], path=db_path, file_name="temporary.docx")

    delete_case_workspace(case.case_id, db_path)
    remaining = {row.review_id for row in list_review_runs(path=db_path)}

    assert case_review_id not in remaining
    assert temp_review_id in remaining


def test_review_run_can_link_case_and_history_can_query(tmp_path: Path):
    db_path = tmp_path / "reviews.sqlite"
    case = create_case_workspace("历史案件", path=db_path)
    review_id = save_review_run([], path=db_path, case_id=case.case_id, file_name="demo.docx", mode_label="单文书字段抽取模式")

    rows = list_review_runs(case.case_id, db_path)
    assert rows[0].review_id == review_id
    assert rows[0].case_id == case.case_id
    assert rows[0].file_name == "demo.docx"


def test_batch_review_failure_does_not_block_successes(tmp_path: Path):
    generate_test_files()
    db_path = tmp_path / "reviews.sqlite"
    case = create_case_workspace("批量案件", path=db_path)
    case_dir = TEST_CASES_DIR / "case_08_case_workspace_batch"
    bad_file = tmp_path / "bad.docx"
    bad_file.write_text("not a docx", encoding="utf-8")
    files = [
        BatchInputFile(path=case_dir / "batch_document_v1.docx", file_name="batch_document_v1.docx", file_size=(case_dir / "batch_document_v1.docx").stat().st_size),
        BatchInputFile(path=bad_file, file_name="bad.docx", file_size=bad_file.stat().st_size),
        BatchInputFile(path=case_dir / "batch_document_clean.docx", file_name="batch_document_clean.docx", file_size=(case_dir / "batch_document_clean.docx").stat().st_size),
    ]

    results = run_batch_review(files, case_dir / "base_info.xlsx", case_dir / "evidence_index.xlsx", case.case_id, db_path=db_path)

    assert sum(result.status == "成功" for result in results) == 2
    assert sum(result.status == "失败" for result in results) == 1
    assert len(list_review_runs(case.case_id, db_path)) == 2


def test_case_package_excel_and_field_differences(tmp_path: Path):
    generate_test_files()
    db_path = tmp_path / "reviews.sqlite"
    case = create_case_workspace("案件包测试", internal_case_no="CASE-08", path=db_path)
    case_dir = TEST_CASES_DIR / "case_08_case_workspace_batch"
    files = [
        BatchInputFile(path=path, file_name=path.name, file_size=path.stat().st_size)
        for path in sorted(case_dir.glob("batch_document*.docx"))
    ]
    results = run_batch_review(files, case_dir / "base_info.xlsx", case_dir / "evidence_index.xlsx", case.case_id, db_path=db_path)
    assert sum(result.status == "成功" for result in results) >= 2

    differences = build_case_field_differences(case.case_id, str(db_path))
    assert differences
    assert any(item.field_type in {"当事人名称", "金额"} for item in differences)

    details = [get_review_run_detail(row.review_id, db_path) for row in list_review_runs(case.case_id, db_path)]
    excel_bytes = build_case_summary_excel(case, [detail for detail in details if detail], differences)
    workbook = load_workbook(filename=Path(_write_bytes(tmp_path / "case_summary.xlsx", excel_bytes)), data_only=True)
    assert {"案件信息", "文书复核汇总", "全部 high 风险", "案件内字段差异"}.issubset(set(workbook.sheetnames))

    package_path = tmp_path / "case_package.zip"
    package_path.write_bytes(build_case_review_package_zip(case.case_id, str(db_path)))
    with ZipFile(package_path) as archive:
        names = set(archive.namelist())
        all_payload = b"".join(archive.read(name) for name in names)
    assert "案件复核汇总.xlsx" in names
    assert "案件复核报告.html" in names
    assert any(name.startswith("单份文书复核包/review_") for name in names)
    assert not any(name.endswith(".sqlite") or name.endswith("review_runs.sqlite") for name in names)
    assert not any(name.endswith(".docx") and "batch_document" in name for name in names)
    assert "2026年2月1日，双方签署《技术服务合同》。".encode("utf-8") not in all_payload


def test_storage_does_not_save_full_document_text(tmp_path: Path):
    generate_test_files()
    db_path = tmp_path / "reviews.sqlite"
    case = create_case_workspace("隐私测试", path=db_path)
    case_dir = TEST_CASES_DIR / "case_08_case_workspace_batch"
    files = [BatchInputFile(path=case_dir / "batch_document_v1.docx", file_name="batch_document_v1.docx", file_size=(case_dir / "batch_document_v1.docx").stat().st_size)]
    run_batch_review(files, case_dir / "base_info.xlsx", case_dir / "evidence_index.xlsx", case.case_id, db_path=db_path)

    conn = sqlite3.connect(db_path)
    try:
        rows = conn.execute("SELECT file_name, findings_json, external_fields_json FROM review_runs").fetchall()
    finally:
        conn.close()
    joined = "\n".join(str(value) for row in rows for value in row)
    assert "2026年2月1日，双方签署《技术服务合同》。" not in joined
    assert "证据1为合同，证据2为验收单" not in joined


def test_single_document_case_has_no_field_differences(tmp_path: Path):
    generate_test_files()
    db_path = tmp_path / "reviews.sqlite"
    case = create_case_workspace("单文书案件", path=db_path)
    case_dir = TEST_CASES_DIR / "case_08_case_workspace_batch"
    file_path = case_dir / "batch_document_v1.docx"
    run_batch_review(
        [BatchInputFile(path=file_path, file_name=file_path.name, file_size=file_path.stat().st_size)],
        case_dir / "base_info.xlsx",
        case_dir / "evidence_index.xlsx",
        case.case_id,
        db_path=db_path,
    )

    assert build_case_field_differences(case.case_id, str(db_path)) == []


def _write_bytes(path: Path, data: bytes) -> str:
    path.write_bytes(data)
    return str(path)

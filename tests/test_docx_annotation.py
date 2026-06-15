from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from app.models.schemas import ExtractedReviewField, ManualReviewRecord, ReviewFinding
from app.reports.docx_annotation import ANNOTATION_DISCLAIMER, annotated_docx_filename, build_annotated_docx
from app.reports.package_export import build_review_package_zip


def _source_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("民事起诉状", level=1)
    document.add_paragraph("原告：上海星澜智能科技有限公司，统一社会信用代码：91310000MA1K000001。")
    document.add_paragraph("诉讼请求：请求支付人民币100,000元。")
    document.save(path)
    return path


def _finding(location: str = "DOCX 段落 2", value: str | None = "91310000MA1K000001") -> ReviewFinding:
    return ReviewFinding(
        risk_id="risk-annotation-1",
        risk_type="SOURCE_MISMATCH",
        severity="high",
        field_name="统一社会信用代码",
        document_value=value,
        reference_value="91310000MA1K000002",
        source_location=location,
        confidence=0.9,
        reason="文书值与基础信息表不一致。",
        suggested_action="建议人工复核统一社会信用代码是否与主体名称一致。",
    )


def _external_field() -> ExtractedReviewField:
    return ExtractedReviewField(
        field_id="field-annotation-1",
        field_type="统一社会信用代码",
        value="9131**********0001",
        role_or_context="原告",
        source_location="DOCX 段落 2",
        confidence=0.9,
        suggested_review_sources=["国家企业信用信息公示系统"],
        suggested_action="建议人工复核统一社会信用代码是否与主体名称一致。",
    )


def _manual_record() -> ManualReviewRecord:
    return ManualReviewRecord(
        record_id="manual-annotation-1",
        review_id="review-annotation-1",
        field_id="field-annotation-1",
        field_type="统一社会信用代码",
        extracted_value="91310000MA1K000001",
        confirmed_value="91310000MA1K000002",
        review_status="需修改",
        review_source="国家企业信用信息公示系统",
        reviewer_note="以客户提供资料为准",
        source_location="DOCX 段落 2",
    )


def _docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs)


def test_annotated_docx_generation_preserves_original_and_adds_prompt(tmp_path: Path):
    source = _source_docx(tmp_path / "source.docx")
    before = hashlib.sha256(source.read_bytes()).hexdigest()

    data = build_annotated_docx(source, [_finding()], [_external_field()], [_manual_record()])

    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    text = _docx_text(data)
    assert "原告：上海星澜智能科技有限公司" in text
    assert "【复核提示】" in text
    assert "风险类型：SOURCE_MISMATCH" in text
    assert "建议人工复核" in text


def test_unlocated_finding_goes_to_summary(tmp_path: Path):
    source = _source_docx(tmp_path / "source.docx")
    finding = _finding(location="文书正文", value="不存在的字段值").model_copy(update={"field_name": "合同编号"})
    data = build_annotated_docx(source, [finding], [], [])
    text = _docx_text(data)

    assert "复核提示汇总" in text
    assert "未定位复核提示" in text
    assert "不存在的字段值" in text


def test_annotated_docx_contains_disclaimer_and_masks_sensitive_values(tmp_path: Path):
    source = _source_docx(tmp_path / "source.docx")
    data = build_annotated_docx(source, [_finding()], [_external_field()], [_manual_record()])
    text = _docx_text(data)

    assert ANNOTATION_DISCLAIMER in text
    assert "91310000MA1K000002" not in text
    assert "9131**********0002" in text


def test_annotated_docx_filename():
    assert annotated_docx_filename("sample_complaint.docx") == "sample_complaint_复核提示版.docx"
    assert annotated_docx_filename("sample.pdf") == "sample_复核提示版.docx"


def test_review_package_can_optionally_include_annotated_docx(tmp_path: Path):
    source = _source_docx(tmp_path / "source.docx")
    data = build_review_package_zip(
        [_finding()],
        [_external_field()],
        [_manual_record()],
        review_id="review-annotation-1",
        source_docx=source,
        source_filename="source.docx",
    )

    with ZipFile(BytesIO(data)) as package:
        names = set(package.namelist())
    assert "文书复核摘要.docx" in names
    assert "复核提示版文书/source_复核提示版.docx" in names


def test_review_package_without_source_docx_excludes_annotated_folder():
    data = build_review_package_zip(
        [_finding()],
        [_external_field()],
        [_manual_record()],
        review_id="review-pdf-like",
    )

    with ZipFile(BytesIO(data)) as package:
        names = set(package.namelist())
    assert "文书复核摘要.docx" in names
    assert not any(name.startswith("复核提示版文书/") for name in names)


def test_annotation_text_avoids_absolute_legal_claims(tmp_path: Path):
    source = _source_docx(tmp_path / "source.docx")
    data = build_annotated_docx(source, [_finding()], [_external_field()], [_manual_record()])
    text = _docx_text(data)

    forbidden = ["系统确认错误", "法律上无效", "自动核验通过", "已完成外部核验", "替代律师判断"]
    assert not any(term in text for term in forbidden)

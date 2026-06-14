from io import BytesIO
from zipfile import ZipFile

from docx import Document

from app.models.schemas import ExtractedReviewField, ManualReviewRecord, ReviewFinding
from app.reports.docx_report import DISCLAIMER, build_docx_report
from app.reports.package_export import build_review_package_zip


def _finding() -> ReviewFinding:
    return ReviewFinding(
        risk_id="risk-1",
        risk_type="FORMAT_ERROR",
        severity="high",
        field_name="身份证号",
        document_value="110101199001011234",
        reference_value=None,
        source_location="段落1",
        confidence=0.9,
        reason="身份证号格式疑似异常，需要复核。",
        suggested_action="建议人工复核该字段格式，并与原始材料核对。",
    )


def _external_field() -> ExtractedReviewField:
    return ExtractedReviewField(
        field_id="field-1",
        field_type="身份证号",
        value="1101**********1234",
        role_or_context="原告",
        source_location="段落1",
        confidence=0.9,
        suggested_review_sources=["身份证件", "内部客户库"],
        suggested_action="建议人工复核身份证号是否与主体身份材料一致。",
    )


def _manual_record() -> ManualReviewRecord:
    return ManualReviewRecord(
        record_id="record-1",
        review_id="review-1",
        field_id="field-1",
        field_type="身份证号",
        extracted_value="110101199001011234",
        confirmed_value="110101199001019999",
        review_status="需修改",
        review_source="身份证件",
        reviewer_note="需与证件原件核对",
        source_location="段落1",
    )


def _docx_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            table_text.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs + table_text)


def test_docx_report_generation_contains_required_sections():
    data = build_docx_report([_finding()], [_external_field()], [_manual_record()], review_id="review-1")
    text = _docx_text(data)
    assert "文书复核摘要" in text
    assert DISCLAIMER in text
    assert "复核发现汇总" in text
    assert "待外部复核字段清单" in text
    assert "人工复核记录" in text


def test_docx_report_generation_without_manual_records():
    data = build_docx_report([_finding()], [_external_field()], [], review_id="review-1")
    text = _docx_text(data)
    assert "暂无人工复核记录" in text


def test_docx_report_masks_sensitive_values():
    data = build_docx_report([_finding()], [_external_field()], [_manual_record()], review_id="review-1")
    text = _docx_text(data)
    assert "110101199001011234" not in text
    assert "110101199001019999" not in text
    assert "****" in text


def test_review_package_zip_contains_excel_html_docx():
    data = build_review_package_zip([_finding()], [_external_field()], [_manual_record()], review_id="review-1")
    with ZipFile(BytesIO(data)) as package:
        names = set(package.namelist())
    assert "复核工作底稿.xlsx" in names
    assert "复核报告.html" in names
    assert "文书复核摘要.docx" in names

from __future__ import annotations

import re
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook

from app.extractors.external_review_fields import build_external_review_fields
from app.parsers.docx_parser import parse_docx
from app.parsers.excel_parser import parse_base_info, parse_evidence_index
from app.parsers.pdf_parser import is_scanned_pdf, parse_pdf
from app.review_engine import review_document, review_single_document
from scripts.generate_test_files import TEST_CASES_DIR, generate_test_files


def test_generate_test_files_creates_expected_cases():
    generated = generate_test_files()
    case_dirs = sorted(path for path in TEST_CASES_DIR.iterdir() if path.is_dir())

    assert generated
    assert TEST_CASES_DIR.is_dir()
    assert len(case_dirs) >= 6
    assert all((case_dir / "README.md").exists() for case_dir in case_dirs)
    assert (TEST_CASES_DIR / "case_08_case_workspace_batch").is_dir()


def test_case_01_full_review_files_and_expected_risks():
    generate_test_files()
    case_dir = TEST_CASES_DIR / "case_01_full_review_basic"

    assert (case_dir / "sample_complaint.docx").exists()
    assert (case_dir / "sample_complaint.pdf").exists()
    assert (case_dir / "base_info.xlsx").exists()
    assert (case_dir / "evidence_index.xlsx").exists()

    chunks = parse_docx(case_dir / "sample_complaint.docx")
    findings = review_document(chunks, parse_base_info(case_dir / "base_info.xlsx"), parse_evidence_index(case_dir / "evidence_index.xlsx"))
    risk_types = {finding.risk_type for finding in findings}

    assert {
        "SOURCE_MISMATCH",
        "AMOUNT_INCONSISTENCY",
        "EVIDENCE_REFERENCE_MISSING",
        "PLACEHOLDER_REMAINING",
        "FORMAT_ERROR",
        "DATE_LOGIC_WARNING",
    }.issubset(risk_types)
    assert not is_scanned_pdf(case_dir / "sample_complaint.pdf")
    assert parse_pdf(case_dir / "sample_complaint.pdf")


def test_case_02_single_document_only_has_no_external_sources():
    generate_test_files()
    case_dir = TEST_CASES_DIR / "case_02_single_document_only"

    assert (case_dir / "single_document.docx").exists()
    assert (case_dir / "single_document.pdf").exists()
    assert not (case_dir / "base_info.xlsx").exists()
    assert not (case_dir / "evidence_index.xlsx").exists()

    chunks = parse_docx(case_dir / "single_document.docx")
    findings = review_single_document(chunks)
    risk_types = {finding.risk_type for finding in findings}
    external_fields = build_external_review_fields(chunks)

    assert "SOURCE_MISMATCH" not in risk_types
    assert "EVIDENCE_REFERENCE_MISSING" not in risk_types
    assert external_fields


def test_case_06_clean_document_has_low_risk_profile():
    generate_test_files()
    case_dir = TEST_CASES_DIR / "case_06_clean_document"
    chunks = parse_docx(case_dir / "clean_document.docx")
    findings = review_document(chunks, parse_base_info(case_dir / "base_info.xlsx"), parse_evidence_index(case_dir / "evidence_index.xlsx"))
    risk_types = {finding.risk_type for finding in findings}

    assert len(findings) <= 1
    assert "SOURCE_MISMATCH" not in risk_types
    assert "EVIDENCE_REFERENCE_MISSING" not in risk_types
    assert "PLACEHOLDER_REMAINING" not in risk_types


def test_generated_test_files_do_not_contain_real_sensitive_patterns():
    generate_test_files()
    text = "\n".join(_extract_text(path) for path in TEST_CASES_DIR.rglob("*") if path.is_file())

    assert not re.search(r"(?<!\d)\d{17}[\dXx](?!\d)", text)
    assert not re.search(r"(?<!\d)1[3-9]\d{9}(?!\d)", text)


def _extract_text(path: Path) -> str:
    if path.suffix.lower() == ".md":
        return path.read_text(encoding="utf-8")
    if path.suffix.lower() == ".docx":
        return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    if path.suffix.lower() == ".pdf":
        with fitz.open(path) as document:
            return "\n".join(page.get_text("text") for page in document)
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, data_only=True)
        values: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values.extend(str(cell) for cell in row if cell is not None)
        return "\n".join(values)
    return ""

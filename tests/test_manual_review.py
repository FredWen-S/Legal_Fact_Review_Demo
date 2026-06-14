from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.manual_review import default_manual_records
from app.models.schemas import ExtractedReviewField, ManualReviewRecord
from app.reports.excel_report import build_excel_report, manual_records_to_dataframe
from app.reports.html_report import build_html_report
from app.storage.sqlite_store import fetch_manual_review_records, save_manual_review_records


def _field() -> ExtractedReviewField:
    return ExtractedReviewField(
        field_id="field-1",
        field_type="身份证号",
        value="1101**********1234",
        role_or_context="原告",
        source_location="段落1",
        confidence=0.9,
        suggested_review_sources=["身份证件", "内部客户库"],
        suggested_action="建议人工复核身份证号是否与主体身份材料一致。",
    )


def test_manual_review_record_schema_default_status():
    record = ManualReviewRecord(
        record_id="r1",
        review_id="review-1",
        field_id="field-1",
        field_type="案号",
        extracted_value="（2024）沪0115民初12345号",
        source_location="段落1",
    )
    assert record.review_status == "未复核"
    assert record.reviewed_at is None


def test_manual_review_record_can_store_confirmed_value():
    record = ManualReviewRecord(
        record_id="r1",
        review_id="review-1",
        field_id="field-1",
        field_type="法定代表人",
        extracted_value="李四",
        confirmed_value="张三",
        review_status="需修改",
        review_source="企查查",
        reviewer_note="以最新公示信息为准",
        source_location="段落1",
        reviewed_at=datetime(2026, 6, 12, 10, 0, 0),
    )
    assert record.confirmed_value == "张三"
    assert record.review_status == "需修改"


def test_sqlite_saves_and_reads_manual_review_records(tmp_path: Path):
    db_path = tmp_path / "reviews.sqlite"
    records = default_manual_records("review-1", [_field()])
    updated = records[0].model_copy(
        update={
            "confirmed_value": "110101199001019999",
            "review_status": "需修改",
            "review_source": "身份证件",
            "reviewer_note": "编号需人工确认",
            "reviewed_at": datetime(2026, 6, 12, 10, 0, 0),
        }
    )
    save_manual_review_records([updated], db_path)
    loaded = fetch_manual_review_records("review-1", db_path)
    assert len(loaded) == 1
    assert loaded[0].confirmed_value == "110101199001019999"
    assert loaded[0].review_source == "身份证件"


def test_excel_report_contains_manual_review_sheet():
    record = ManualReviewRecord(
        record_id="r1",
        review_id="review-1",
        field_id="field-1",
        field_type="案号",
        extracted_value="（2024）沪0115民初12345号",
        review_status="已确认",
        review_source="法院系统",
        source_location="段落1",
        reviewed_at=datetime(2026, 6, 12, 10, 0, 0),
    )
    workbook = load_workbook(BytesIO(build_excel_report([], [_field()], [record])))
    assert "人工复核记录" in workbook.sheetnames


def test_html_report_contains_manual_review_section():
    record = ManualReviewRecord(
        record_id="r1",
        review_id="review-1",
        field_id="field-1",
        field_type="案号",
        extracted_value="（2024）沪0115民初12345号",
        review_status="未复核",
        source_location="段落1",
    )
    html = build_html_report([], [_field()], [record], mode_label="单文书字段抽取模式")
    assert "人工复核记录" in html
    assert "未复核" in html


def test_manual_review_export_masks_sensitive_values():
    record = ManualReviewRecord(
        record_id="r1",
        review_id="review-1",
        field_id="field-1",
        field_type="身份证号",
        extracted_value="110101199001011234",
        confirmed_value="110101199001019999",
        review_status="需修改",
        source_location="段落1",
    )
    text = manual_records_to_dataframe([record]).to_string()
    assert "110101199001011234" not in text
    assert "110101199001019999" not in text
    assert "****" in text

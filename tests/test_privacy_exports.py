from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook

from app.models.schemas import ExtractedReviewField, ManualReviewRecord, ReviewFinding
from app.privacy.policy import PrivacyPolicy
from app.reports.docx_report import build_docx_report
from app.reports.excel_report import build_excel_report
from app.reports.html_report import build_html_report
from app.reports.package_export import build_review_package_zip


SENSITIVE_TEXT = "Contact test@example.com or 13812345678 with ID 110101199001011234"


def _finding() -> ReviewFinding:
    return ReviewFinding(
        risk_id="privacy-1",
        risk_type="FORMAT_ERROR",
        severity="medium",
        field_name="联系信息",
        document_value=SENSITIVE_TEXT,
        reference_value=None,
        source_location="段落1",
        confidence=0.8,
        reason="联系方式需要人工复核。",
        suggested_action="建议人工复核该字段。",
    )


def _external_field() -> ExtractedReviewField:
    return ExtractedReviewField(
        field_id="field-1",
        field_type="联系信息",
        value=SENSITIVE_TEXT,
        source_location="段落1",
        confidence=0.8,
        suggested_review_sources=["synthetic source"],
        suggested_action="建议人工复核该字段。",
    )


def _manual_record() -> ManualReviewRecord:
    return ManualReviewRecord(
        record_id="manual-1",
        review_id="review-1",
        field_id="field-1",
        field_type="联系信息",
        extracted_value=SENSITIVE_TEXT,
        review_status="未复核",
        source_location="段落1",
    )


def _docx_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            table_text.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs + table_text)


def test_html_excel_word_and_zip_include_privacy_summary():
    html = build_html_report([_finding()], [_external_field()], [_manual_record()])
    assert "Privacy Summary / 隐私摘要" in html
    assert "This is a synthetic local demo." in html

    workbook = load_workbook(BytesIO(build_excel_report([_finding()])), data_only=True)
    assert "Privacy Summary" in workbook.sheetnames
    privacy_values = [cell.value for row in workbook["Privacy Summary"].iter_rows() for cell in row]
    assert "This is a synthetic local demo." in privacy_values

    docx_text = _docx_text(build_docx_report([_finding()]))
    assert "Privacy Summary" in docx_text
    assert "This is a synthetic local demo." in docx_text

    with ZipFile(BytesIO(build_review_package_zip([_finding()]))) as package:
        names = set(package.namelist())
        privacy_json = package.read("privacy_summary.json").decode("utf-8")
    assert "privacy_summary.json" in names
    assert "This is a synthetic local demo." in privacy_json


def test_default_report_policy_does_not_forcibly_redact_issue_text():
    html = build_html_report([_finding()])

    assert "test@example.com" in html
    assert "13812345678" in html


def test_report_redaction_policy_masks_sensitive_issue_text():
    policy = PrivacyPolicy(redact_report_sensitive_fields=True)

    html = build_html_report([_finding()], privacy_policy=policy)
    workbook = load_workbook(BytesIO(build_excel_report([_finding()], privacy_policy=policy)), data_only=True)
    docx_text = _docx_text(build_docx_report([_finding()], privacy_policy=policy))

    assert "test@example.com" not in html
    assert "13812345678" not in html
    assert "110101199001011234" not in html
    assert "te***@example.com" in html

    issues_values = [cell.value for row in workbook["Issues"].iter_rows() for cell in row]
    joined_excel = "\n".join(str(value) for value in issues_values)
    assert "test@example.com" not in joined_excel
    assert "13812345678" not in joined_excel

    assert "test@example.com" not in docx_text
    assert "13812345678" not in docx_text

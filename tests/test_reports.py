from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from app.models.schemas import (
    ExtractedReviewField,
    KnowledgeContext,
    KnowledgeMatch,
    KnowledgeRetrievalStatus,
    ManualReviewRecord,
    ReviewFinding,
    StructuredEntity,
    StructuredExtractionResult,
)
from app.reports.docx_report import build_docx_report
from app.reports.excel_report import build_excel_report, findings_to_dataframe
from app.reports.html_report import build_html_report


def _finding() -> ReviewFinding:
    return ReviewFinding(
        risk_id="r1",
        risk_type="PLACEHOLDER_REMAINING",
        severity="high",
        field_name="文书中的占位符",
        document_value="待补充",
        reference_value=None,
        source_location="段落1",
        confidence=0.95,
        reason="A placeholder marker is still present in the document.",
        suggested_action="建议人工复核该处内容。",
    )


def _external_field() -> ExtractedReviewField:
    return ExtractedReviewField(
        field_id="f1",
        field_type="案号",
        value="（2024）沪0115民初12345号",
        role_or_context="案件信息",
        source_location="段落1",
        confidence=0.9,
        suggested_review_sources=["法院系统", "案件台账"],
        suggested_action="建议人工复核案号格式和案件信息是否正确。",
    )


def _manual_record() -> ManualReviewRecord:
    return ManualReviewRecord(
        record_id="r1",
        review_id="review-1",
        field_id="f1",
        field_type="案号",
        extracted_value="（2024）沪0115民初12345号",
        review_status="未复核",
        source_location="段落1",
    )


def _structured_extraction() -> StructuredExtractionResult:
    return StructuredExtractionResult(
        mode="mock",
        model_name="deterministic-mock-extractor",
        fallback_reason="no_api_key",
        entities=[
            StructuredEntity(
                entity_type="credit_code",
                value="91310000MA1K000001",
                source_location="段落1",
            )
        ],
    )


def _knowledge_context() -> KnowledgeContext:
    return KnowledgeContext(
        client_profile=KnowledgeMatch(
            kb_type="client",
            doc_id="client-demo",
            title="合成客户档案",
            score=4.2,
            min_score=3.0,
            path="kb/clients/client_demo.json",
            snippet="Synthetic profile",
            match_reason="score 4.2 >= min_score 3.0",
        ),
        retrieval_statuses=[
            KnowledgeRetrievalStatus(
                kb_type="client",
                status="matched",
                min_score=3.0,
                best_score=4.2,
                best_doc_id="client-demo",
                match_reason="Best client profile score 4.2 met minimum 3.0.",
            )
        ],
    )


def _docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return f"{text}\n{table_text}"


def test_issue_dataframe_uses_new_issue_schema():
    df = findings_to_dataframe([_finding()])
    assert list(df.columns[:6]) == [
        "Severity",
        "Issue Type",
        "Original Text",
        "Suggested Text",
        "Evidence Source",
        "Explanation",
    ]
    assert df.iloc[0]["Severity"] == "HIGH"


def test_excel_report_generation():
    data = build_excel_report([_finding()], [_external_field()], [_manual_record()])
    assert data[:2] == b"PK"


def test_html_report_generation():
    html = build_html_report([_finding()], [_external_field()], [_manual_record()])
    assert "LLM + RAG Legal Document Review Report" in html
    assert "Manual Review Notes / 人工复核记录" in html
    assert "Review Workbench / 待外部复核字段" in html
    assert "Structured Extraction / 结构化抽取" in html


def test_reports_include_config_version_and_document_profile():
    excel = build_excel_report([_finding()], config_version="case09-custom-v1", document_profile="lawyer_letter")
    workbook = load_workbook(BytesIO(excel), data_only=True)
    meta = workbook["Metadata"]
    values = {(row[0].value, row[1].value) for row in meta.iter_rows(min_row=2, max_col=2)}
    assert ("Config Version", "case09-custom-v1") in values
    assert ("Document Profile", "lawyer_letter") in values

    html = build_html_report([_finding()], config_version="case09-custom-v1", document_profile="lawyer_letter")
    assert "case09-custom-v1" in html
    assert "lawyer_letter" in html

    docx = build_docx_report([_finding()], config_version="case09-custom-v1", document_profile="lawyer_letter")
    combined = _docx_text(docx)
    assert "case09-custom-v1" in combined
    assert "lawyer_letter" in combined


def test_docx_report_includes_ai_rag_context_section():
    docx = build_docx_report(
        [_finding()],
        [_external_field()],
        [_manual_record()],
        structured_extraction=_structured_extraction(),
        knowledge_context=_knowledge_context(),
    )
    text = _docx_text(docx)

    assert "AI / RAG Context" in text
    assert "Fallback Reason" in text
    assert "no_api_key" in text
    assert "score 4.2 >= min_score 3.0" in text


def test_docx_report_ai_rag_context_handles_missing_context():
    text = _docx_text(build_docx_report([_finding()]))

    assert "AI / RAG Context" in text
    assert "Structured extraction context is not available." in text
    assert "Retrieved knowledge context is not available." in text
